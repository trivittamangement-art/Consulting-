"""Gera o relatório mensal do PPR Comex a partir da planilha de controle.

Uso:
    python3 relatorio_mensal.py <planilha.xlsx> --mes 2026-09 [--preencher] [--saida pasta]

Saídas (em ppr-comex/relatorios/AAAA-MM/ ou --saida):
    Relatorio_PPR_AAAA-MM.docx / .md    relatório das 3 metas
    Descritivo_Embalagens_AAAA-MM.docx  sugestões de uso de embalagens, por pedido, para a produção
    Pendencias_AAAA-MM.md               o que falta o gerente informar

--preencher grava na aba Acompanhamento uma linha para cada cliente que não
recomprou no mês e ainda não tem motivo, para o gerente completar.

O script lê só as colunas de entrada (não depende das fórmulas estarem
calculadas), então funciona com planilha recém-editada em qualquer programa.
"""
import argparse
import datetime as dt
from pathlib import Path

import pandas as pd
from dateutil.relativedelta import relativedelta
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook

import estrutura as E

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]


# ------------------------------------------------------------------ leitura

def ler(xlsx, aba, colunas, chave):
    try:
        df = pd.read_excel(xlsx, sheet_name=aba)
    except ValueError:
        return pd.DataFrame(columns=colunas)
    df = df[[c for c in colunas if c in df.columns]].copy()
    for c in colunas:
        if c not in df.columns:
            df[c] = None
    df = df[df[chave].notna() & (df[chave].astype(str).str.strip() != "")]
    return df.reset_index(drop=True)


def datas(df, *cols):
    for c in cols:
        df[c] = pd.to_datetime(df[c], errors="coerce", dayfirst=True)
    return df


def num(df, *cols):
    for c in cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def parametros(xlsx):
    p = pd.read_excel(xlsx, sheet_name="Parametros", usecols="A:B", header=0).dropna()
    d = dict(zip(p.iloc[:, 0], p.iloc[:, 1]))
    sla = pd.read_excel(xlsx, sheet_name="Parametros", usecols="E:F", header=0).dropna()
    return d, dict(zip(sla.iloc[:, 0], sla.iloc[:, 1]))


# ------------------------------------------------------------------ documento

class Doc:
    """Monta o relatório uma vez e exporta em Markdown e Word."""

    def __init__(self):
        self.itens = []

    def h(self, txt, nivel=1):
        self.itens.append(("h", nivel, txt))

    def p(self, txt):
        self.itens.append(("p", txt))

    def lista(self, linhas):
        if linhas:
            self.itens.append(("l", list(linhas)))

    def tabela(self, df):
        if df is not None and len(df):
            self.itens.append(("t", df.fillna("").astype(str)))

    def md(self):
        out = []
        for it in self.itens:
            if it[0] == "h":
                out.append("#" * it[1] + " " + it[2])
            elif it[0] == "p":
                out.append(it[1])
            elif it[0] == "l":
                out.append("\n".join("- " + x for x in it[1]))
            else:
                df = it[1]
                out.append("| " + " | ".join(df.columns) + " |")
                out[-1] += "\n|" + "---|" * len(df.columns)
                for row in df.itertuples(index=False):
                    out[-1] += "\n| " + " | ".join(str(v).replace("|", "/") for v in row) + " |"
            out.append("")
        return "\n".join(out)

    def docx(self, caminho):
        d = Document()
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.5)
        st = d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        for it in self.itens:
            if it[0] == "h":
                hd = d.add_heading(it[2], level=it[1])
                for r in hd.runs:
                    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            elif it[0] == "p":
                d.add_paragraph(it[1])
            elif it[0] == "l":
                for x in it[1]:
                    d.add_paragraph(x, style="List Bullet")
            else:
                df = it[1]
                t = d.add_table(rows=1, cols=len(df.columns))
                t.style = "Light Grid Accent 1"
                for j, c in enumerate(df.columns):
                    t.rows[0].cells[j].text = str(c)
                for row in df.itertuples(index=False):
                    cells = t.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = str(v)
                for row in t.rows:
                    for c in row.cells:
                        for par in c.paragraphs:
                            for r in par.runs:
                                r.font.size = Pt(9)
                d.add_paragraph()
        d.save(caminho)


def fmt_n(v, casas=1):
    if v is None or pd.isna(v):
        return "-"
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_rs(v):
    return "R$ " + fmt_n(v, 2)


def fmt_pct(v):
    return "-" if v is None or pd.isna(v) else f"{v * 100:.0f}%".replace(".", ",")


def fmt_d(v):
    return "" if v is None or pd.isna(v) else pd.Timestamp(v).strftime("%d/%m/%Y")


def ok(cond):
    return "Atingida" if cond else "Abaixo da meta"


# ------------------------------------------------------------------ meta 2

def recompra_do_mes(cli, ped, ini, fim):
    """Classifica cada cliente no mês pela periodicidade cadastrada."""
    linhas = []
    for _, c in cli.iterrows():
        cod = c["Cód. Cliente"]
        pc = ped[ped["Cód. Cliente"] == cod]
        antes = pc[pc["Data do pedido"] < ini]
        no_mes = pc[(pc["Data do pedido"] >= ini) & (pc["Data do pedido"] <= fim)]
        per = c["Periodicidade (meses)"]
        ultimo = antes["Data do pedido"].max() if len(antes) else pd.NaT
        prevista = ultimo + relativedelta(months=int(per)) if pd.notna(ultimo) and pd.notna(per) else pd.NaT
        if len(no_mes):
            if pd.isna(ultimo) or ultimo < ini - relativedelta(months=12):
                sit = "Novo/reativado"
            elif pd.notna(prevista) and prevista > fim:
                sit = "Antecipou"
            else:
                sit = "Recomprou"
        elif pd.notna(prevista) and prevista <= fim:
            sit = "Atrasado" if prevista < ini else "Não recomprou"
        else:
            sit = "Fora do ciclo"
        if str(c.get("Status cadastro") or "").strip() in ("Inativo", "Perdido") and sit in ("Atrasado", "Não recomprou"):
            sit = "Perdido"
        linhas.append({
            "Cód.": cod, "Cliente": c["Cliente"], "País": c["País"],
            "Periodicidade": "" if pd.isna(per) else f"{int(per)} m",
            "Último pedido": fmt_d(ultimo), "Prevista": fmt_d(prevista), "_prevista": prevista,
            "Situação": sit, "Pedidos no mês": len(no_mes),
            "t no mês": no_mes["Toneladas"].sum(),
        })
    return pd.DataFrame(linhas, columns=["Cód.", "Cliente", "País", "Periodicidade", "Último pedido", "Prevista",
                                         "_prevista", "Situação", "Pedidos no mês", "t no mês"])


# ------------------------------------------------------------------ principal

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("--mes", required=True, help="AAAA-MM")
    ap.add_argument("--preencher", action="store_true")
    ap.add_argument("--saida")
    a = ap.parse_args()

    xlsx = Path(a.planilha)
    ini = pd.Timestamp(a.mes + "-01")
    fim = ini + relativedelta(months=1) - pd.Timedelta(days=1)
    nome_mes = f"{MESES[ini.month - 1]}/{ini.year}"
    hoje = pd.Timestamp(dt.date.today())

    par, sla = parametros(xlsx)
    ano_base = int(par.get("Ano base", ini.year - 1))
    ano_meta = int(par.get("Ano meta", ini.year))
    meta_pos = float(par.get("Meta de aumento de positivação", 0.20))
    meta_vol = float(par.get("Meta de volume (t/mês)", 75))
    prazo_padrao = int(par.get("Prazo de embarque padrão (dias)", 30))
    sla = sla or E.SLA_PADRAO

    cli = num(ler(xlsx, "Clientes", E.CLIENTES, "Cód. Cliente"), "Periodicidade (meses)")
    ped = num(datas(ler(xlsx, "Pedidos", E.PEDIDOS, "Nº Pedido"), "Data do pedido",
                    "Embarque negociado", "Embarque realizado"), "Toneladas", "Valor (USD)")
    ped = ped[ped["Data do pedido"].notna()]
    acp = datas(ler(xlsx, "Acompanhamento", E.ACOMPANHAMENTO, "Cód. Cliente"), "Mês ref.", "Data de retorno")
    atv = num(datas(ler(xlsx, "S_Atividades", E.ATIVIDADES, "Nº Pedido"), "Data solicitação",
                    "Data conclusão", "Última cobrança"), "Prazo (dias)", "Nº cobranças")
    cus = num(datas(ler(xlsx, "Custos", E.CUSTOS, "Descrição da ação"), "Mês ref."),
              "Custo de referência (R$)", "Custo realizado (R$)")
    emb = num(datas(ler(xlsx, "Embalagens", E.EMBALAGENS, "Código"), "Data de entrada"),
              "Qtd em estoque", "Custo unit. (R$)")
    uso = num(datas(ler(xlsx, "Uso_Embalagens", E.USO_EMBALAGENS, "Código embalagem"), "Data"), "Qtd utilizada")

    nome_cli = dict(zip(cli["Cód. Cliente"], cli["Cliente"]))
    idioma_cli = dict(zip(cli["Cód. Cliente"], cli["Idioma"]))
    ped["Cliente"] = ped["Cód. Cliente"].map(nome_cli)
    ped.loc[ped["Embarque negociado"].isna(), "Embarque negociado"] = \
        ped["Data do pedido"] + pd.Timedelta(days=prazo_padrao)

    pend = []  # o que falta o gerente informar
    doc = Doc()
    doc.h(f"PPR Comex · Relatório de {nome_mes}")
    doc.p(f"Período apurado: {fmt_d(ini)} a {fmt_d(fim)}. Gerado em {fmt_d(hoje)} a partir da planilha "
          f"{xlsx.name}. Ano base da positivação: {ano_base}.")

    # ---------------- meta 1
    pm = ped[(ped["Data do pedido"] >= ini) & (ped["Data do pedido"] <= fim)]
    ton_mes = pm["Toneladas"].sum()
    ano = ped[(ped["Data do pedido"].dt.year == ano_meta) & (ped["Data do pedido"] <= fim)]
    meses_ano = ini.month if ini.year == ano_meta else 12
    media = ano["Toneladas"].sum() / meses_ano if meses_ano else 0

    # ---------------- meta 2
    base = ped[ped["Data do pedido"].dt.year == ano_base]
    cli_base, ped_base = base["Cód. Cliente"].nunique(), len(base)
    lim_mesmo = fim - relativedelta(years=ano_meta - ano_base)
    base_ytd = base[base["Data do pedido"] <= lim_mesmo]
    meta_ytd = ano
    cli_ytd, ped_ytd = meta_ytd["Cód. Cliente"].nunique(), len(meta_ytd)
    cli_base_ytd, ped_base_ytd = base_ytd["Cód. Cliente"].nunique(), len(base_ytd)
    cresc_cli = cli_ytd / cli_base_ytd - 1 if cli_base_ytd else None
    cresc_ped = ped_ytd / ped_base_ytd - 1 if ped_base_ytd else None
    rec = recompra_do_mes(cli, ped, ini, fim)
    previstos = rec[rec["Situação"].isin(["Recomprou", "Não recomprou", "Atrasado", "Perdido"])]
    recomp = (rec["Situação"] == "Recomprou").sum()
    retencao = recomp / len(previstos) if len(previstos) else None
    sem = rec[rec["Situação"].isin(["Não recomprou", "Atrasado", "Perdido"])].copy()
    acp_mes = acp[acp["Mês ref."].dt.to_period("M") == ini.to_period("M")] if len(acp) else acp
    mot = dict(zip(acp_mes["Cód. Cliente"], acp_mes.to_dict("records"))) if len(acp_mes) else {}
    sem["Motivo"] = sem["Cód."].map(lambda c: (mot.get(c) or {}).get("Motivo") or "PENDENTE")
    sem["Detalhe"] = sem["Cód."].map(lambda c: (mot.get(c) or {}).get("Detalhe do motivo") or "")
    sem["Ação"] = sem["Cód."].map(lambda c: (mot.get(c) or {}).get("Ação combinada") or "")
    sem["Retorno"] = sem["Cód."].map(lambda c: fmt_d((mot.get(c) or {}).get("Data de retorno")))
    sem["Chance"] = sem["Cód."].map(lambda c: (mot.get(c) or {}).get("Chance de recompra") or "")
    faltam_motivo = sem[sem["Motivo"] == "PENDENTE"]

    # ---------------- meta 3a
    atv["Cliente"] = atv["Cliente"].where(atv["Cliente"].notna() & (atv["Cliente"].astype(str) != ""),
                                          atv["Nº Pedido"].map(dict(zip(ped["Nº Pedido"], ped["Cliente"]))))
    atv["Prazo"] = atv["Prazo (dias)"].fillna(atv["Departamento"].map(sla)).fillna(0)
    atv["Limite"] = atv["Data solicitação"] + pd.to_timedelta(atv["Prazo"], unit="D")
    atv["Status"] = atv.apply(
        lambda r: ("No prazo" if r["Data conclusão"] <= r["Limite"] else "Com atraso") if pd.notna(r["Data conclusão"])
        else ("Atrasada" if hoje > r["Limite"] else "Em aberto"), axis=1) if len(atv) else []
    atv["Atraso"] = atv.apply(
        lambda r: max(0, ((r["Data conclusão"] if pd.notna(r["Data conclusão"]) else hoje) - r["Limite"]).days),
        axis=1) if len(atv) else []
    # atividades do mês: solicitadas no mês, concluídas no mês ou ainda abertas
    am = atv[((atv["Data solicitação"] >= ini) & (atv["Data solicitação"] <= fim)) |
             ((atv["Data conclusão"] >= ini) & (atv["Data conclusão"] <= fim)) |
             (atv["Data conclusão"].isna())] if len(atv) else atv
    emb_mes = ped[(ped["Embarque realizado"] >= ini) & (ped["Embarque realizado"] <= fim)]
    no_prazo_emb = (emb_mes["Embarque realizado"] <= emb_mes["Embarque negociado"]).mean() if len(emb_mes) else None

    # ---------------- meta 3b
    cm = cus[cus["Mês ref."].dt.to_period("M") == ini.to_period("M")].copy() if len(cus) else cus
    if len(cm):
        cm["Economia"] = cm["Custo de referência (R$)"].fillna(0) - cm["Custo realizado (R$)"].fillna(0)
    eco_mes = cm["Economia"].sum() if len(cm) else 0
    ca = cus[(cus["Mês ref."].dt.year == ano_meta) & (cus["Mês ref."] <= fim)] if len(cus) else cus
    eco_ano = (ca["Custo de referência (R$)"].fillna(0) - ca["Custo realizado (R$)"].fillna(0)).sum() if len(ca) else 0

    # ---------------- meta 3c
    if len(emb):
        emb["Valor"] = emb["Qtd em estoque"].fillna(0) * emb["Custo unit. (R$)"].fillna(0)
        emb["Idade (m)"] = emb["Data de entrada"].map(
            lambda d: (hoje.year - d.year) * 12 + hoje.month - d.month if pd.notna(d) else None)
        idade_obs = int(par.get("Idade para obsoleta (meses)", 12))
        emb["Obsoleta"] = (emb["Status"].astype(str) == "Obsoleta") | (emb["Idade (m)"].fillna(0) >= idade_obs) & \
                          (emb["Status"].astype(str) != "Ativa")
    obs = emb[emb["Obsoleta"]] if len(emb) else emb
    um = uso[(uso["Data"] >= ini) & (uso["Data"] <= fim)].copy() if len(uso) else uso
    custo_emb = dict(zip(emb["Código"], emb["Custo unit. (R$)"])) if len(emb) else {}
    if len(um):
        um["Economia"] = um["Qtd utilizada"].fillna(0) * um["Código embalagem"].map(custo_emb).fillna(0)
    valor_uso_mes = um["Economia"].sum() if len(um) else 0

    # ================================================================ resumo
    doc.h("1. Resumo das metas", 2)
    resumo = pd.DataFrame([
        ["Meta 1 · Volume médio", f"{fmt_n(meta_vol, 0)} t/mês",
         f"{fmt_n(ton_mes)} t no mês · média {fmt_n(media)} t/mês no ano", ok(media >= meta_vol)],
        ["Meta 2 · Positivação +20%", f"{fmt_pct(meta_pos)} sobre {ano_base}",
         f"Clientes {fmt_pct(cresc_cli)} · pedidos {fmt_pct(cresc_ped)} (acumulado x mesmo período de {ano_base})",
         ok(cresc_cli is not None and cresc_cli >= meta_pos)],
        ["Meta 2 · Manutenção (recompra)", "Clientes previstos recomprando",
         f"{recomp} de {len(previstos)} ({fmt_pct(retencao)}) · {len(sem)} sem recompra, "
         f"{len(sem) - len(faltam_motivo)} com motivo registrado", "Ver seção 3"],
        ["Meta 3a · S.A. e prazo", f"Embarque em {prazo_padrao} dias",
         f"{(am['Status'] == 'No prazo').sum() if len(am) else 0} de {am['Data conclusão'].notna().sum() if len(am) else 0} "
         f"atividades concluídas no prazo · embarques no prazo: {fmt_pct(no_prazo_emb)}", "Ver seção 4"],
        ["Meta 3b · Redução de custo", "Apresentar ações do mês",
         f"{len(cm)} ações · {fmt_rs(eco_mes)} no mês · {fmt_rs(eco_ano)} no ano", "Entregue" if len(cm) else "Pendente"],
        ["Meta 3c · Giro de embalagens", "Mapear e propor uso",
         f"{fmt_rs(obs['Valor'].sum() if len(obs) else 0)} em obsoletas · {fmt_rs(valor_uso_mes)} reaproveitado no mês",
         "Entregue" if len(um) or len(emb) else "Pendente"],
    ], columns=["Meta", "Referência", "Resultado", "Status"])
    doc.tabela(resumo)

    # ================================================================ meta 1
    doc.h("2. Meta 1 · Volume", 2)
    doc.p(f"Toneladas com pedido lançado em {nome_mes}: {fmt_n(ton_mes)} t em {len(pm)} pedidos. "
          f"Média mensal de {ano_meta} até o mês: {fmt_n(media)} t (meta {fmt_n(meta_vol, 0)} t).")
    if len(pm):
        top = pm.groupby("Cliente", dropna=False)["Toneladas"].sum().sort_values(ascending=False).head(10)
        doc.tabela(pd.DataFrame({"Cliente": top.index, "Toneladas": [fmt_n(v) for v in top.values]}))
    if media < meta_vol and ini.month < 12:
        resta = 12 - ini.month
        nec = (meta_vol * 12 - ano["Toneladas"].sum()) / resta if resta else 0
        doc.p(f"Para fechar o ano na média, os próximos {resta} meses precisam de {fmt_n(nec)} t/mês.")

    # ================================================================ meta 2
    doc.h("3. Meta 2 · Positivação e manutenção de clientes", 2)
    doc.tabela(pd.DataFrame([
        ["Clientes positivados", cli_base, cli_base_ytd, cli_ytd, fmt_pct(cresc_cli),
         int(-(-cli_base * (1 + meta_pos) // 1))],
        ["Pedidos", ped_base, ped_base_ytd, ped_ytd, fmt_pct(cresc_ped), int(-(-ped_base * (1 + meta_pos) // 1))],
    ], columns=["Indicador", f"{ano_base} (ano)", f"{ano_base} até {lim_mesmo.strftime('%d/%m')}",
                f"{ano_meta} até {fim.strftime('%d/%m')}", "Variação", f"Meta {ano_meta} (ano)"]))
    novos = rec[rec["Situação"] == "Novo/reativado"]
    contagem = rec["Situação"].value_counts()
    doc.p("Situação dos clientes no mês, pela periodicidade cadastrada: " + "; ".join(
        f"{k}: {v}" for k, v in contagem.items()) + ".")
    if len(novos):
        doc.p("Clientes novos ou reativados no mês (somam para os +20%): " +
              ", ".join(novos["Cliente"].astype(str)) + ".")
    doc.h("3.1 Clientes que não recompraram no mês", 3)
    if len(sem):
        doc.tabela(sem[["Cliente", "País", "Periodicidade", "Último pedido", "Prevista", "Situação",
                        "Motivo", "Detalhe", "Ação", "Retorno"]])
        com = sem[sem["Motivo"] != "PENDENTE"]
        if len(com):
            doc.h("3.2 Motivos de não recompra", 3)
            mc = com["Motivo"].value_counts()
            doc.tabela(pd.DataFrame({"Motivo": mc.index, "Clientes": mc.values,
                                     "%": [fmt_pct(v / len(com)) for v in mc.values]}))
            doc.p(f"Principal causa no mês: {mc.index[0].lower()} ({mc.iloc[0]} de {len(com)} clientes).")
        if len(faltam_motivo):
            doc.p(f"{len(faltam_motivo)} cliente(s) ainda sem motivo registrado (PENDENTE).")
            pend.append("Meta 2: informar motivo, ação e data de retorno de: " +
                        ", ".join(faltam_motivo["Cliente"].astype(str)))
    else:
        doc.p("Todos os clientes previstos para o mês recompraram.")
    if not len(cli) or not len(base):
        pend.append(f"Meta 2: carregar a base de clientes e os pedidos de {ano_base} (abas Clientes e Pedidos).")
    sem_per = cli[cli["Periodicidade (meses)"].isna()]
    if len(sem_per):
        pend.append(f"Meta 2: {len(sem_per)} cliente(s) sem periodicidade cadastrada: " +
                    ", ".join(sem_per["Cliente"].astype(str).head(15)))

    # ================================================================ meta 3a
    doc.h("4. Meta 3a · Gestão das solicitações de atividade (S.A.)", 2)
    if len(am):
        concl = am[am["Data conclusão"].notna()]
        doc.p(f"Atividades acompanhadas no mês: {len(am)} em {am['Nº Pedido'].nunique()} pedidos. "
              f"Concluídas: {len(concl)} ({(concl['Status'] == 'No prazo').sum()} no prazo). "
              f"Em aberto: {am['Data conclusão'].isna().sum()} "
              f"({(am['Status'] == 'Atrasada').sum()} já fora do prazo). "
              f"Cobranças registradas: {int(am['Nº cobranças'].fillna(0).sum())}.")
        g = am.groupby("Departamento").agg(
            Atividades=("Nº Pedido", "count"),
            Concluídas=("Data conclusão", "count"),
            No_prazo=("Status", lambda s: (s == "No prazo").sum()),
            Atrasadas=("Status", lambda s: s.isin(["Com atraso", "Atrasada"]).sum()),
            Atraso_médio=("Atraso", lambda s: s[s > 0].mean() if (s > 0).any() else 0),
            Cobranças=("Nº cobranças", lambda s: int(s.fillna(0).sum())),
        ).reset_index()
        g["% no prazo"] = [fmt_pct(n / c) if c else "-" for n, c in zip(g["No_prazo"], g["Concluídas"])]
        g["Atraso_médio"] = g["Atraso_médio"].map(lambda v: fmt_n(v, 1) + " d")
        g = g.rename(columns={"No_prazo": "No prazo", "Atraso_médio": "Atraso médio"})
        g["SLA (dias)"] = g["Departamento"].map(sla).fillna("")
        doc.tabela(g[["Departamento", "SLA (dias)", "Atividades", "Concluídas", "No prazo", "% no prazo",
                      "Atrasadas", "Atraso médio", "Cobranças"]])
        pior = g.sort_values("Atrasadas", ascending=False).iloc[0]
        if pior["Atrasadas"]:
            doc.p(f"Setor com mais atrasos: {pior['Departamento']} ({pior['Atrasadas']} atividades).")
        causas = am["Causa do atraso"].dropna().astype(str).str.strip()
        causas = causas[causas != ""].value_counts()
        if len(causas):
            doc.p("Causas de atraso registradas: " + "; ".join(f"{k} ({v})" for k, v in causas.items()) + ".")
        abertas = am[am["Status"] == "Atrasada"].copy()
        if len(abertas):
            doc.h("4.1 Atividades atrasadas em aberto (cobrar)", 3)
            abertas["Limite"] = abertas["Limite"].map(fmt_d)
            abertas["Última cobrança"] = abertas["Última cobrança"].map(fmt_d)
            doc.tabela(abertas[["Nº Pedido", "Cliente", "Departamento", "Atividade", "Limite", "Atraso",
                                "Nº cobranças", "Última cobrança"]].rename(columns={"Atraso": "Dias de atraso"}))
        sem_cob = am[(am["Status"].isin(["Atrasada", "Com atraso"])) & (am["Nº cobranças"].fillna(0) == 0)]
        if len(sem_cob):
            pend.append(f"Meta 3a: {len(sem_cob)} atividade(s) atrasada(s) sem cobrança registrada; "
                        "confirmar se houve cobrança (data e meio).")
    else:
        doc.p("Nenhuma atividade registrada no mês.")
        pend.append("Meta 3a: enviar a lista de S.A. dos pedidos do mês (pedido, setor, atividade, data de "
                    "solicitação, data de conclusão, cobranças).")
    # risco de embarque
    abertos = ped[ped["Embarque realizado"].isna() & (ped["Data do pedido"] <= fim)].copy()
    if len(abertos) and len(atv):
        pend_atv = atv[atv["Data conclusão"].isna()].groupby("Nº Pedido").size()
        abertos["Atividades pendentes"] = abertos["Nº Pedido"].map(pend_atv).fillna(0).astype(int)
        abertos["Dias p/ embarque"] = (abertos["Embarque negociado"] - hoje).dt.days
        risco = abertos[(abertos["Atividades pendentes"] > 0) & (abertos["Dias p/ embarque"] <= 10)]
        if len(risco):
            doc.h("4.2 Pedidos com risco de perder o prazo de embarque", 3)
            risco = risco.sort_values("Dias p/ embarque")
            risco["Embarque negociado"] = risco["Embarque negociado"].map(fmt_d)
            doc.tabela(risco[["Nº Pedido", "Cliente", "Embarque negociado", "Dias p/ embarque",
                              "Atividades pendentes"]])
    if len(emb_mes):
        doc.p(f"Embarques realizados no mês: {len(emb_mes)}; no prazo negociado: {fmt_pct(no_prazo_emb)}. "
              f"Tempo médio pedido-embarque: "
              f"{fmt_n((emb_mes['Embarque realizado'] - emb_mes['Data do pedido']).dt.days.mean(), 0)} dias.")
        fora = emb_mes[emb_mes["Embarque realizado"] > emb_mes["Embarque negociado"]]
        if len(fora):
            doc.p("Embarques fora do prazo: " + ", ".join(
                f"{r['Nº Pedido']} ({r['Cliente']}, +{(r['Embarque realizado'] - r['Embarque negociado']).days} d)"
                for _, r in fora.iterrows()) + ".")

    # ================================================================ meta 3b
    doc.h("5. Meta 3b · Redução de custos", 2)
    if len(cm):
        t = cm.copy()
        t["Referência"] = t["Custo de referência (R$)"].map(fmt_rs)
        t["Realizado"] = t["Custo realizado (R$)"].map(fmt_rs)
        t["Economia"] = t["Economia"].map(fmt_rs)
        doc.tabela(t[["Categoria", "Descrição da ação", "Tipo de ganho", "Referência", "Realizado",
                      "Economia", "Evidência"]])
        cat = cm.groupby("Categoria")["Economia"].sum().sort_values(ascending=False)
        doc.p(f"Total no mês: {fmt_rs(eco_mes)}. Acumulado no ano: {fmt_rs(eco_ano)}. Maior contribuição: "
              f"{cat.index[0]} ({fmt_rs(cat.iloc[0])}).")
        sem_ev = cm[cm["Evidência"].isna() | (cm["Evidência"].astype(str).str.strip() == "")]
        if len(sem_ev):
            pend.append(f"Meta 3b: {len(sem_ev)} ação(ões) de custo sem evidência (e-mail, cotação, nota).")
    else:
        doc.p("Nenhuma ação de custo registrada no mês.")
        pend.append("Meta 3b: listar o que aconteceu no mês com amostras, inspeções, viagens, impressões, "
                    "courier e fretes (quem pagou, quanto custaria x quanto custou).")
    if valor_uso_mes:
        doc.p(f"Embalagens reaproveitadas no mês evitaram {fmt_rs(valor_uso_mes)} em compra de embalagem "
              "nova (registrado também na meta 3c).")

    # ================================================================ meta 3c
    doc.h("6. Meta 3c · Giro de embalagens", 2)
    if len(emb):
        doc.p(f"Itens mapeados: {len(emb)}. Obsoletos: {len(obs)}, somando "
              f"{fmt_n(obs['Qtd em estoque'].sum(), 0)} unidades e {fmt_rs(obs['Valor'].sum())} parados.")
        top = obs.sort_values("Valor", ascending=False).head(10).copy()
        top["Valor"] = top["Valor"].map(fmt_rs)
        top["Qtd em estoque"] = top["Qtd em estoque"].map(lambda v: fmt_n(v, 0))
        doc.tabela(top[["Código", "Descrição", "Idioma", "Qtd em estoque", "Valor", "Idade (m)"]])
    else:
        pend.append("Meta 3c: enviar o inventário de embalagens (código, descrição, idioma, quantidade, "
                    "custo unitário, data de entrada).")
    if len(um):
        t = um.copy()
        t["Descrição"] = t["Código embalagem"].map(dict(zip(emb["Código"], emb["Descrição"])))
        t["Economia"] = t["Economia"].map(fmt_rs)
        t["Data"] = t["Data"].map(fmt_d)
        doc.h("6.1 Reaproveitado no mês", 3)
        doc.tabela(t[["Data", "Nº Pedido", "Código embalagem", "Descrição", "Qtd utilizada", "Forma de uso",
                      "Aprovado pela produção", "Economia"]])

    # sugestões por pedido em aberto
    sug = []
    if len(obs):
        candidatos = ped[(ped["Embarque realizado"].isna()) | (ped["Embarque realizado"] > fim)]
        candidatos = candidatos[candidatos["Data do pedido"] >= ini - relativedelta(months=2)]
        for _, pr in candidatos.iterrows():
            idi = str(idioma_cli.get(pr["Cód. Cliente"], "") or "").lower()
            for _, e in obs.iterrows():
                ei = str(e["Idioma"] or "").lower()
                tipo = str(e["Tipo"] or "")
                linha = str(e["Produto/linha"] or "").lower()
                mesmo_prod = linha in ("", "geral", "nan") or linha in str(pr["Produto"] or "").lower()
                mesma_lingua = ei and ei != "neutra" and idi and (ei in idi or idi in ei)
                forma = None
                if mesma_lingua and mesmo_prod:
                    forma = "Mesmo idioma"
                elif mesma_lingua and tipo == "Caixa master":
                    forma = "Caixa externa/reforço ou bonificação (produto diferente)"
                elif tipo == "Divisória/reforço" or ei == "neutra":
                    forma = "Reforço/acessório"
                if forma:
                    sug.append({"Nº Pedido": pr["Nº Pedido"], "Cliente": pr["Cliente"],
                                "Idioma do cliente": idioma_cli.get(pr["Cód. Cliente"], ""),
                                "Produto do pedido": pr["Produto"], "Código": e["Código"],
                                "Embalagem": e["Descrição"], "Saldo": fmt_n(e["Qtd em estoque"], 0),
                                "Uso sugerido": forma})
    sug = pd.DataFrame(sug)
    if len(sug):
        ordem = {"Mesmo idioma": 0, "Reforço/acessório": 1}
        sug["_o"] = sug["Uso sugerido"].map(ordem).fillna(2)
        sug = sug.sort_values(["Nº Pedido", "_o"]).groupby("Nº Pedido").head(3).drop(columns="_o")
        doc.h("6.2 Oportunidades nos pedidos em carteira", 3)
        doc.tabela(sug[["Nº Pedido", "Cliente", "Código", "Embalagem", "Uso sugerido"]].head(30))
        doc.p("As sugestões dependem de validação da produção e da qualidade (produto compatível, medida da caixa, "
              "rotulagem exigida pelo país). O descritivo por pedido segue em arquivo separado.")

    # ================================================================ próximos passos
    doc.h("7. Próximos passos", 2)
    passos = []
    if len(sem):
        passos.append(f"Contatar os {len(sem)} clientes sem recompra, começando pelos atrasados há mais tempo.")
    if len(am) and (am["Status"] == "Atrasada").any():
        passos.append("Cobrar as atividades atrasadas da seção 4.1 e registrar cada cobrança na planilha.")
    if len(sug):
        passos.append(f"Validar com a produção as {len(sug)} sugestões de uso de embalagem.")
    if not len(cm):
        passos.append("Registrar ao menos uma ação de custo (amostra paga pelo cliente, inspeção conjunta, "
                      "documento digital).")
    doc.lista(passos)

    # ================================================================ saída
    pasta = Path(a.saida) if a.saida else Path(__file__).resolve().parent.parent / "relatorios" / a.mes
    pasta.mkdir(parents=True, exist_ok=True)
    (pasta / f"Relatorio_PPR_{a.mes}.md").write_text(doc.md(), encoding="utf-8")
    doc.docx(pasta / f"Relatorio_PPR_{a.mes}.docx")

    if len(sug):
        d = Document()
        d.add_heading(f"Descritivo de embalagens para a produção · {nome_mes}", 1)
        d.add_paragraph("Proposta de uso de embalagens em estoque nos pedidos abaixo. Peço que a produção confirme "
                        "a viabilidade de cada item (medida, compatibilidade com o produto e exigência de rotulagem "
                        "do destino) e devolva com Sim ou Não.")
        for pedido, grp in sug.groupby("Nº Pedido"):
            r0 = grp.iloc[0]
            d.add_heading(f"Pedido {pedido} · {r0['Cliente']}", 2)
            d.add_paragraph(f"Idioma do cliente: {r0['Idioma do cliente']}. Produto: {r0['Produto do pedido']}.")
            for _, s in grp.iterrows():
                d.add_paragraph(f"{s['Código']} · {s['Embalagem']} (saldo {s['Saldo']} un.): {s['Uso sugerido']}. "
                                "Quantidade a usar: ____  Aprovado: ( ) Sim ( ) Não", style="List Bullet")
        rod = d.add_paragraph("Após a aprovação, o uso será lançado na aba Uso_Embalagens da planilha de controle.")
        rod.alignment = WD_ALIGN_PARAGRAPH.LEFT
        d.save(pasta / f"Descritivo_Embalagens_{a.mes}.docx")

    p = [f"# Pendências para fechar o relatório de {nome_mes}", ""]
    p += [f"- {x}" for x in pend] or ["- Nenhuma pendência."]
    (pasta / f"Pendencias_{a.mes}.md").write_text("\n".join(p) + "\n", encoding="utf-8")

    if a.preencher and len(faltam_motivo):
        wb = load_workbook(xlsx)
        ws = wb["Acompanhamento"]
        r = 2
        while ws.cell(row=r, column=2).value not in (None, ""):
            r += 1
        for _, s in faltam_motivo.iterrows():
            ws.cell(row=r, column=1, value=ini.to_pydatetime())
            ws.cell(row=r, column=2, value=s["Cód."])
            ws.cell(row=r, column=3, value=s["Cliente"])
            ws.cell(row=r, column=4, value=s["_prevista"].to_pydatetime() if pd.notna(s["_prevista"]) else None)
            ws.cell(row=r, column=5, value=s["Situação"])
            ws.cell(row=r, column=11, value="Aberta")
            r += 1
        wb.save(xlsx)
        print(f"{len(faltam_motivo)} linha(s) incluída(s) na aba Acompanhamento.")

    print(pasta)
    for x in pend:
        print("PENDENTE:", x)


if __name__ == "__main__":
    main()
