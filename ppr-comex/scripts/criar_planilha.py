"""Gera a planilha de controle do PPR Comex.

Uso:
    python3 criar_planilha.py                    # modelo em branco
    python3 criar_planilha.py --exemplo          # modelo com dados fictícios para teste
    python3 criar_planilha.py --saida arquivo.xlsx
"""
import argparse
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import estrutura as E

AZUL = "1F3864"
CINZA = "F2F2F2"
CALC = "E2EFDA"  # colunas calculadas: verde claro, não editar
VERMELHO = "F8CBAD"
VERDE = "C6EFCE"
AMARELO = "FFEB9C"

LINHAS = {"Clientes": 500, "Pedidos": 3000, "Acompanhamento": 2000,
          "S_Atividades": 2000, "Custos": 300, "Embalagens": 500,
          "Uso_Embalagens": 1000}

REF = "Parametros!$B$2"
ANO_BASE = "Parametros!$B$3"
ANO_META = "Parametros!$B$4"
FIM_MES = f"EOMONTH({REF},0)"

fino = Side(style="thin", color="BFBFBF")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)


def cabecalho(ws, entradas, calculadas=()):
    for i, nome in enumerate(list(entradas) + [c[0] for c in calculadas], start=1):
        cel = ws.cell(row=1, column=i, value=nome)
        calc = i > len(entradas)
        cel.font = Font(bold=True, color="FFFFFF" if not calc else "1F1F1F")
        cel.fill = PatternFill("solid", fgColor=AZUL if not calc else "A9D08E")
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cel.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(34, len(nome) + 4))
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "B2" if ws.title in ("Clientes", "Embalagens") else "A2"
    ultima = get_column_letter(len(entradas) + len(calculadas))
    ws.auto_filter.ref = f"A1:{ultima}1"


def formulas(ws, entradas, calculadas, n):
    """calculadas: lista de (nome, template com {r}, formato)."""
    inicio = len(entradas) + 1
    for j, (_, tpl, fmt) in enumerate(calculadas):
        col = inicio + j
        for r in range(2, n + 2):
            c = ws.cell(row=r, column=col, value="=" + tpl.format(r=r))
            c.fill = PatternFill("solid", fgColor=CALC)
            if fmt:
                c.number_format = fmt


def formato(ws, col, fmt, n):
    for r in range(2, n + 2):
        ws.cell(row=r, column=col).number_format = fmt


def lista(ws, col, nome_lista, n):
    itens = E.LISTAS[nome_lista]
    idx = list(E.LISTAS).index(nome_lista) + 1
    letra = get_column_letter(idx)
    dv = DataValidation(type="list", formula1=f"=Listas!${letra}$2:${letra}${len(itens) + 1}",
                        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    cl = get_column_letter(col)
    dv.add(f"{cl}2:{cl}{n + 1}")


def destaque(ws, col, n, vermelhos=(), verdes=(), amarelos=()):
    cl = get_column_letter(col)
    rng = f"{cl}2:{cl}{n + 1}"
    for textos, cor in ((vermelhos, VERMELHO), (verdes, VERDE), (amarelos, AMARELO)):
        for t in textos:
            ws.conditional_formatting.add(
                rng, FormulaRule(formula=[f'{cl}2="{t}"'], fill=PatternFill("solid", fgColor=cor)))


def aba_leiame(wb):
    ws = wb.active
    ws.title = "LEIA-ME"
    linhas = [
        ("PPR Comex · Planilha de controle", True),
        ("", False),
        ("Colunas azuis: você (ou o agente) preenche. Colunas verdes: fórmulas, não editar.", False),
        ("Comece pela aba Parametros: troque o Mês de referência todo mês (sempre dia 01).", False),
        ("", False),
        ("Meta 1 · Volume (75 t/mês): aba Pedidos (só lançar pedidos; o Painel calcula).", True),
        ("Meta 2 · Positivação +20%: abas Clientes, Pedidos e Acompanhamento.", True),
        ("  Clientes: um por linha, com a periodicidade de compra (1, 2, 3, 4, 6 ou 12 meses).", False),
        ("  Pedidos: lançar os pedidos do ano base (2025) e do ano meta (2026).", False),
        ("  Acompanhamento: um registro por cliente que não recomprou no mês, com motivo e ação.", False),
        ("Meta 3a · Gestão das S.A.: aba S_Atividades (uma linha por atividade solicitada).", True),
        ("  O prazo de cada setor vem da tabela de SLA da aba Parametros se a coluna Prazo ficar vazia.", False),
        ("  Registre cada cobrança feita (Nº cobranças e Última cobrança): é a prova da sua gestão.", False),
        ("Meta 3b · Redução de custo: aba Custos (uma linha por ação; referência x realizado).", True),
        ("Meta 3c · Giro de embalagens: abas Embalagens (estoque) e Uso_Embalagens (o que foi usado).", True),
        ("", False),
        ("Relatório do mês: python3 ppr-comex/scripts/relatorio_mensal.py <planilha> --mes AAAA-MM", False),
    ]
    for i, (txt, negrito) in enumerate(linhas, start=1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(bold=negrito, size=14 if i == 1 else 11, color=AZUL if negrito else "000000")
    ws.column_dimensions["A"].width = 110


def aba_parametros(wb, mes_ref):
    ws = wb.create_sheet("Parametros")
    for j, h in enumerate(["Parâmetro", "Valor", "Observação"], start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    for i, (nome, valor, obs) in enumerate(E.PARAMETROS, start=2):
        ws.cell(row=i, column=1, value=nome)
        if nome == "Mês de referência":
            valor = mes_ref
        c = ws.cell(row=i, column=2, value=valor)
        ws.cell(row=i, column=3, value=obs)
    ws["B2"].number_format = "mmm/yyyy"
    ws["B5"].number_format = "0%"
    for j, h in enumerate(["Departamento", "Prazo padrão (dias)"], start=5):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    for i, (dep, dias) in enumerate(E.SLA_PADRAO.items(), start=2):
        ws.cell(row=i, column=5, value=dep)
        ws.cell(row=i, column=6, value=dias)
    for col, w in zip("ABCDEF", (34, 14, 46, 3, 28, 20)):
        ws.column_dimensions[col].width = w


def aba_listas(wb):
    ws = wb.create_sheet("Listas")
    for j, (nome, itens) in enumerate(E.LISTAS.items(), start=1):
        ws.cell(row=1, column=j, value=nome).font = Font(bold=True)
        for i, item in enumerate(itens, start=2):
            ws.cell(row=i, column=j, value=item)
        ws.column_dimensions[get_column_letter(j)].width = 30


def aba_clientes(wb):
    ws = wb.create_sheet("Clientes")
    n = LINHAS["Clientes"]
    calc = [
        ("Pedidos ano base", f'IF(A{{r}}="","",COUNTIFS(Pedidos!$C:$C,A{{r}},Pedidos!$K:$K,{ANO_BASE}))', "0"),
        ("Toneladas ano base", f'IF(A{{r}}="","",SUMIFS(Pedidos!$E:$E,Pedidos!$C:$C,A{{r}},Pedidos!$K:$K,{ANO_BASE}))', "0.0"),
        ("Pedidos ano meta", f'IF(A{{r}}="","",COUNTIFS(Pedidos!$C:$C,A{{r}},Pedidos!$K:$K,{ANO_META}))', "0"),
        ("Último pedido", f'IF(A{{r}}="","",IF(COUNTIFS(Pedidos!$C:$C,A{{r}},Pedidos!$B:$B,"<="&{FIM_MES})=0,"",_xlfn.MAXIFS(Pedidos!$B:$B,Pedidos!$C:$C,A{{r}},Pedidos!$B:$B,"<="&{FIM_MES})))', "dd/mm/yyyy"),
        ("Próxima compra prevista", 'IF(OR(L{r}="",F{r}=""),"",EDATE(L{r},F{r}))', "dd/mm/yyyy"),
        ("Situação no mês", f'IF(A{{r}}="","",IF(L{{r}}="","Sem histórico",IF(AND(L{{r}}>={REF},L{{r}}<={FIM_MES}),"Comprou no mês",IF(M{{r}}="","Sem periodicidade",IF(M{{r}}<{REF},"Atrasado",IF(M{{r}}<={FIM_MES},"Previsto no mês","Em dia"))))))', None),
    ]
    cabecalho(ws, E.CLIENTES, calc)
    formulas(ws, E.CLIENTES, calc, n)
    lista(ws, 6, "Periodicidade", n)
    lista(ws, 7, "Status cadastro", n)
    destaque(ws, 14, n, vermelhos=["Atrasado"], verdes=["Comprou no mês"], amarelos=["Previsto no mês"])
    ws.column_dimensions["B"].width = 30
    return ws


def aba_pedidos(wb):
    ws = wb.create_sheet("Pedidos")
    n = LINHAS["Pedidos"]
    calc = [
        ("Cliente", 'IF(C{r}="","",IFERROR(INDEX(Clientes!$B:$B,MATCH(C{r},Clientes!$A:$A,0)),"Cód. não cadastrado"))', None),
        ("Ano", 'IF(B{r}="","",YEAR(B{r}))', "0"),
        ("Mês", 'IF(B{r}="","",MONTH(B{r}))', "0"),
        ("Dias até embarque", 'IF(OR(B{r}="",I{r}=""),"",I{r}-B{r})', "0"),
        ("Prazo de embarque", 'IF(B{r}="","",IF(I{r}="",IF(H{r}="","",IF(TODAY()>H{r},"Atrasado (em aberto)","Em aberto")),IF(I{r}<=H{r},"No prazo","Fora do prazo")))', None),
    ]
    cabecalho(ws, E.PEDIDOS, calc)
    formulas(ws, E.PEDIDOS, calc, n)
    for col in (2, 8, 9):
        formato(ws, col, "dd/mm/yyyy", n)
    formato(ws, 5, "0.00", n)
    formato(ws, 6, "#,##0.00", n)
    destaque(ws, 14, n, vermelhos=["Fora do prazo", "Atrasado (em aberto)"], verdes=["No prazo"])
    return ws


def aba_acompanhamento(wb):
    ws = wb.create_sheet("Acompanhamento")
    n = LINHAS["Acompanhamento"]
    cabecalho(ws, E.ACOMPANHAMENTO)
    formato(ws, 1, "mmm/yyyy", n)
    formato(ws, 4, "dd/mm/yyyy", n)
    formato(ws, 10, "dd/mm/yyyy", n)
    lista(ws, 5, "Situação", n)
    lista(ws, 6, "Motivo", n)
    lista(ws, 11, "Status da ação", n)
    lista(ws, 12, "Chance de recompra", n)
    for col, w in ((3, 30), (6, 30), (7, 40), (8, 40)):
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws


def aba_atividades(wb):
    ws = wb.create_sheet("S_Atividades")
    n = LINHAS["S_Atividades"]
    calc = [
        ("Data limite", 'IF(E{r}="","",E{r}+IF(F{r}="",IFERROR(VLOOKUP(C{r},Parametros!$E:$F,2,0),0),F{r}))', "dd/mm/yyyy"),
        ("Status", 'IF(E{r}="","",IF(G{r}="",IF(TODAY()>L{r},"Atrasada","Em aberto"),IF(G{r}<=L{r},"No prazo","Com atraso")))', None),
        ("Dias de atraso", 'IF(E{r}="","",MAX(0,IF(G{r}="",TODAY(),G{r})-L{r}))', "0"),
    ]
    cabecalho(ws, E.ATIVIDADES, calc)
    formulas(ws, E.ATIVIDADES, calc, n)
    for col in (5, 7, 9):
        formato(ws, col, "dd/mm/yyyy", n)
    lista(ws, 3, "Departamento", n)
    destaque(ws, 13, n, vermelhos=["Atrasada", "Com atraso"], verdes=["No prazo"])
    ws.column_dimensions["D"].width = 36
    return ws


def aba_custos(wb):
    ws = wb.create_sheet("Custos")
    n = LINHAS["Custos"]
    calc = [("Economia (R$)", 'IF(AND(E{r}="",F{r}=""),"",E{r}-F{r})', "#,##0.00")]
    cabecalho(ws, E.CUSTOS, calc)
    formulas(ws, E.CUSTOS, calc, n)
    formato(ws, 1, "mmm/yyyy", n)
    formato(ws, 5, "#,##0.00", n)
    formato(ws, 6, "#,##0.00", n)
    lista(ws, 2, "Categoria custo", n)
    lista(ws, 4, "Tipo de ganho", n)
    ws.column_dimensions["C"].width = 50
    return ws


def aba_embalagens(wb):
    ws = wb.create_sheet("Embalagens")
    n = LINHAS["Embalagens"]
    calc = [
        ("Valor em estoque (R$)", 'IF(A{r}="","",H{r}*J{r})', "#,##0.00"),
        ("Idade (meses)", 'IF(K{r}="","",DATEDIF(K{r},TODAY(),"m"))', "0"),
        ("Qtd já utilizada", 'IF(A{r}="","",SUMIFS(Uso_Embalagens!$E:$E,Uso_Embalagens!$D:$D,A{r}))', "#,##0"),
    ]
    cabecalho(ws, E.EMBALAGENS, calc)
    formulas(ws, E.EMBALAGENS, calc, n)
    formato(ws, 8, "#,##0", n)
    formato(ws, 10, "#,##0.00", n)
    formato(ws, 11, "dd/mm/yyyy", n)
    lista(ws, 3, "Tipo embalagem", n)
    lista(ws, 12, "Status embalagem", n)
    ws.column_dimensions["B"].width = 36
    return ws


def aba_uso(wb):
    ws = wb.create_sheet("Uso_Embalagens")
    n = LINHAS["Uso_Embalagens"]
    calc = [
        ("Descrição", 'IF(D{r}="","",IFERROR(INDEX(Embalagens!$B:$B,MATCH(D{r},Embalagens!$A:$A,0)),"Código não cadastrado"))', None),
        ("Economia (R$)", 'IF(D{r}="","",E{r}*IFERROR(INDEX(Embalagens!$J:$J,MATCH(D{r},Embalagens!$A:$A,0)),0))', "#,##0.00"),
    ]
    cabecalho(ws, E.USO_EMBALAGENS, calc)
    formulas(ws, E.USO_EMBALAGENS, calc, n)
    formato(ws, 1, "dd/mm/yyyy", n)
    lista(ws, 6, "Forma de uso", n)
    lista(ws, 7, "Sim/Não", n)
    return ws


def aba_painel(wb):
    ws = wb.create_sheet("Painel", 1)
    ini, fim = REF, FIM_MES
    na = 'Pedidos!$B:$B,">="&{0},Pedidos!$B:$B,"<="&{1}'.format(ini, fim)
    blocos = [
        ("META 1 · VOLUME", None, None),
        ("Toneladas vendidas no mês", f"SUMIFS(Pedidos!$E:$E,{na})", "0.0"),
        ("Média mensal no ano (até o mês de referência)",
         f'SUMIFS(Pedidos!$E:$E,Pedidos!$K:$K,{ANO_META},Pedidos!$B:$B,"<="&{fim})/MONTH({ini})', "0.0"),
        ("Meta (t/mês)", "Parametros!$B$6", "0"),
        ("META 2 · POSITIVAÇÃO", None, None),
        ("Clientes positivados no ano base", 'COUNTIF(Clientes!$I:$I,">0")', "0"),
        ("Meta de clientes no ano (+20%)", "ROUNDUP(B7*(1+Parametros!$B$5),0)", "0"),
        ("Clientes positivados no ano meta (acumulado)", 'COUNTIF(Clientes!$K:$K,">0")', "0"),
        ("Pedidos no ano base", f"COUNTIF(Pedidos!$K:$K,{ANO_BASE})", "0"),
        ("Meta de pedidos no ano (+20%)", "ROUNDUP(B10*(1+Parametros!$B$5),0)", "0"),
        ("Pedidos no ano meta (acumulado)", f"COUNTIF(Pedidos!$K:$K,{ANO_META})", "0"),
        ("Clientes que compraram no mês", 'COUNTIF(Clientes!$N:$N,"Comprou no mês")', "0"),
        ("Clientes previstos no mês e ainda sem pedido", 'COUNTIF(Clientes!$N:$N,"Previsto no mês")', "0"),
        ("Clientes atrasados na recompra", 'COUNTIF(Clientes!$N:$N,"Atrasado")', "0"),
        ("Motivos registrados no mês", f"COUNTIFS(Acompanhamento!$A:$A,{ini})", "0"),
        ("META 3a · S.A. E PRAZO DE EMBARQUE", None, None),
        ("Atividades concluídas no mês",
         f'COUNTIFS(S_Atividades!$G:$G,">="&{ini},S_Atividades!$G:$G,"<="&{fim})', "0"),
        ("% concluídas no prazo",
         f'IFERROR(COUNTIFS(S_Atividades!$G:$G,">="&{ini},S_Atividades!$G:$G,"<="&{fim},S_Atividades!$M:$M,"No prazo")/B18,"")', "0%"),
        ("Atividades em aberto atrasadas (hoje)", 'COUNTIF(S_Atividades!$M:$M,"Atrasada")', "0"),
        ("Cobranças feitas (atividades do mês)",
         f'SUMIFS(S_Atividades!$H:$H,S_Atividades!$E:$E,">="&{ini},S_Atividades!$E:$E,"<="&{fim})', "0"),
        ("Pedidos embarcados no mês",
         f'COUNTIFS(Pedidos!$I:$I,">="&{ini},Pedidos!$I:$I,"<="&{fim})', "0"),
        ("% embarcados no prazo negociado",
         f'IFERROR(COUNTIFS(Pedidos!$I:$I,">="&{ini},Pedidos!$I:$I,"<="&{fim},Pedidos!$N:$N,"No prazo")/B22,"")', "0%"),
        ("META 3b · REDUÇÃO DE CUSTO", None, None),
        ("Economia registrada no mês (R$)", f"SUMIFS(Custos!$I:$I,Custos!$A:$A,{ini})", "#,##0.00"),
        ("Economia acumulada no ano (R$)",
         f'SUMIFS(Custos!$I:$I,Custos!$A:$A,">="&DATE({ANO_META},1,1),Custos!$A:$A,"<="&{fim})', "#,##0.00"),
        ("META 3c · GIRO DE EMBALAGENS", None, None),
        ("Valor parado em embalagens obsoletas (R$)", 'SUMIFS(Embalagens!$O:$O,Embalagens!$L:$L,"Obsoleta")', "#,##0.00"),
        ("Valor reaproveitado no mês (R$)",
         f'SUMIFS(Uso_Embalagens!$J:$J,Uso_Embalagens!$A:$A,">="&{ini},Uso_Embalagens!$A:$A,"<="&{fim})', "#,##0.00"),
        ("Unidades reaproveitadas no mês",
         f'SUMIFS(Uso_Embalagens!$E:$E,Uso_Embalagens!$A:$A,">="&{ini},Uso_Embalagens!$A:$A,"<="&{fim})', "#,##0"),
    ]
    ws["A1"] = "Painel PPR Comex · mês de referência:"
    ws["B1"] = f"={REF}"
    ws["B1"].number_format = "mmm/yyyy"
    ws["A1"].font = ws["B1"].font = Font(bold=True, size=13, color=AZUL)
    for i, (rotulo, f, fmt) in enumerate(blocos, start=2):
        a = ws.cell(row=i, column=1, value=rotulo)
        if f is None:
            a.font = Font(bold=True, color="FFFFFF")
            a.fill = PatternFill("solid", fgColor=AZUL)
            ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=AZUL)
            continue
        b = ws.cell(row=i, column=2, value="=" + f)
        b.number_format = fmt
        b.fill = PatternFill("solid", fgColor=CINZA)
        a.border = b.border = BORDA
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18


# ---------------------------------------------------------------- exemplo

def dados_exemplo(wb):
    """Dados fictícios só para testar fórmulas e o relatório."""
    rnd = random.Random(7)
    paises = [("Paraguai", "Espanhol"), ("Chile", "Espanhol"), ("Uruguai", "Espanhol"),
              ("Angola", "Português"), ("Emirados Árabes", "Inglês/Árabe"),
              ("Estados Unidos", "Inglês"), ("Japão", "Japonês"), ("Bolívia", "Espanhol"),
              ("Peru", "Espanhol"), ("Gana", "Inglês")]
    produtos = ["Produto A 1kg", "Produto B 5kg", "Produto C granel 25kg"]
    ws = wb["Clientes"]
    clientes = []
    for i in range(1, 27):
        pais, idioma = rnd.choice(paises)
        per = rnd.choice([1, 2, 2, 3, 3, 4, 6])
        cod = f"C{i:03d}"
        nome = f"Cliente Exemplo {i:02d}"
        clientes.append((cod, nome, per, i > 22))  # os 4 últimos são novos em 2026
        for j, v in enumerate([cod, nome, pais, idioma, rnd.choice(["Direto", "Agente LatAm", "Trading"]),
                               per, "Ativo", ""], start=1):
            ws.cell(row=i + 1, column=j, value=v)
    ws = wb["Pedidos"]
    r, num = 2, 1000
    for cod, _, per, novo in clientes:
        anos = [2026] if novo else [2025, 2026]
        for ano in anos:
            mes = rnd.randint(1, per) if not novo else rnd.randint(5, 9)
            while mes <= (12 if ano == 2025 else 9):
                if ano == 2026 and rnd.random() < 0.18:  # cliente que pulou a compra
                    mes += per
                    continue
                data = dt.date(ano, mes, rnd.randint(1, 26))
                neg = data + dt.timedelta(days=30)
                atraso = rnd.choice([-6, -3, 0, 0, 2, 5, 9])
                real = neg + dt.timedelta(days=atraso)
                if real > dt.date(2026, 9, 30):
                    real = None
                num += 1
                for j, v in enumerate([f"PV{num}", data, cod, rnd.choice(produtos),
                                       round(rnd.uniform(3, 9), 1), round(rnd.uniform(15000, 60000), 2),
                                       rnd.choice(["FOB", "CFR", "CIF"]), neg, real], start=1):
                    ws.cell(row=r, column=j, value=v)
                r += 1
                mes += per
    cli_ped = {wb["Pedidos"].cell(row=i, column=1).value:
               dict((c[0], c[1]) for c in clientes).get(wb["Pedidos"].cell(row=i, column=3).value, "")
               for i in range(2, r)}
    ws = wb["S_Atividades"]
    deps = list(E.SLA_PADRAO)
    atividades = {"Compras/Suprimentos": "Compra de embalagem/insumos", "Marketing/Arte": "Arte de rótulo no idioma",
                  "Qualidade/Regulatório": "Certificado sanitário / análise", "PCP/Produção": "Programar produção",
                  "Logística/Embarque": "Booking e coleta do contêiner", "Financeiro": "Liberação de crédito / carta de crédito",
                  "Documentação/Despachante": "Documentos de exportação", "Comercial/Comex": "Proforma e confirmação"}
    r = 2
    for k in range(12):
        ped = f"PV{num - k}"
        inicio = dt.date(2026, 9, rnd.randint(1, 20))
        for dep in rnd.sample(deps, 5):
            sol = inicio + dt.timedelta(days=rnd.randint(0, 3))
            sla = E.SLA_PADRAO[dep]
            dur = sla + rnd.choice([-3, -1, 0, 0, 2, 4, 8])
            concl = sol + dt.timedelta(days=max(1, dur))
            if concl > dt.date(2026, 10, 5):
                concl = None
            cob = rnd.choice([0, 1, 1, 2, 3])
            causa = rnd.choice(["Fornecedor atrasou", "Aguardando aprovação do cliente", "Fila da produção", ""]) \
                if dur > sla else ""
            linha = [ped, cli_ped.get(ped, ""), dep, atividades[dep], sol, None, concl, cob,
                     sol + dt.timedelta(days=sla) if cob else None, causa, ""]
            for j, v in enumerate(linha, start=1):
                ws.cell(row=r, column=j, value=v)
            r += 1
    ws = wb["Custos"]
    mes = dt.date(2026, 9, 1)
    custos = [
        (mes, "Amostras", "Cliente C004 pagou o courier das amostras (antes por nossa conta)", "Custo transferido ao cliente", 850, 0, "E-mail 12/09", "Gerente Comex"),
        (mes, "Inspeções", "Inspeção de dois pedidos (C002 e C009) feita na mesma visita", "Custo evitado", 3200, 1600, "Agenda da certificadora", "Gerente Comex"),
        (mes, "Impressões / documentos", "Documentos de embarque enviados só em PDF", "Economia direta", 420, 90, "Controle de impressões", "Gerente Comex"),
        (mes, "Viagens internacionais", "Visita ao cliente C011 substituída por reunião on-line", "Custo evitado", 14500, 0, "Orçamento da agência", "Gerente Comex"),
    ]
    for i, linha in enumerate(custos, start=2):
        for j, v in enumerate(linha, start=1):
            ws.cell(row=i, column=j, value=v)
    ws = wb["Embalagens"]
    emb = [
        ("EMB-101", "Caixa master Produto A - espanhol (layout 2021)", "Caixa master", "Produto A", "Espanhol", "Paraguai", "40x30x25", 3200, "un", 3.80, dt.date(2021, 5, 10), "Obsoleta", "Galpão 2", ""),
        ("EMB-102", "Cartucho Produto A - inglês (EUA, 2020)", "Cartucho", "Produto A", "Inglês", "Estados Unidos", "15x10x5", 18000, "un", 0.65, dt.date(2020, 11, 3), "Obsoleta", "Galpão 2", "Tabela nutricional antiga"),
        ("EMB-103", "Caixa master Produto B - árabe/inglês", "Caixa master", "Produto B", "Inglês/Árabe", "Emirados Árabes", "50x35x30", 1500, "un", 5.20, dt.date(2022, 2, 14), "Obsoleta", "Galpão 1", ""),
        ("EMB-104", "Caixa master neutra sem impressão", "Caixa master", "Geral", "Neutra", "", "40x30x25", 900, "un", 3.10, dt.date(2023, 7, 1), "Ativa", "Galpão 1", ""),
        ("EMB-105", "Saco 25kg - japonês", "Saco/filme", "Produto C", "Japonês", "Japão", "", 2400, "un", 2.10, dt.date(2021, 9, 20), "Obsoleta", "Galpão 3", ""),
        ("EMB-106", "Divisória papelão ondulado", "Divisória/reforço", "Geral", "Neutra", "", "38x28", 6000, "un", 0.40, dt.date(2022, 4, 4), "Obsoleta", "Galpão 3", ""),
        ("EMB-107", "Caixa master Produto B - português (Angola 2019)", "Caixa master", "Produto B", "Português", "Angola", "50x35x30", 2100, "un", 5.00, dt.date(2019, 8, 30), "Obsoleta", "Galpão 2", "Logo antigo"),
    ]
    for i, linha in enumerate(emb, start=2):
        for j, v in enumerate(linha, start=1):
            ws.cell(row=i, column=j, value=v)
    ws = wb["Uso_Embalagens"]
    usos = [(dt.date(2026, 9, 15), f"PV{num - 3}", "", "EMB-106", 400, "Reforço/acessório", "Sim", "Divisória entre camadas")]
    for i, linha in enumerate(usos, start=2):
        for j, v in enumerate(linha, start=1):
            ws.cell(row=i, column=j, value=v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--exemplo", action="store_true")
    ap.add_argument("--saida")
    ap.add_argument("--mes", help="mês de referência AAAA-MM")
    a = ap.parse_args()
    base = Path(__file__).resolve().parent.parent / "modelos"
    saida = Path(a.saida) if a.saida else base / ("PPR_Comex_EXEMPLO.xlsx" if a.exemplo else "PPR_Comex_Controle.xlsx")
    if a.mes:
        mes_ref = dt.datetime.strptime(a.mes + "-01", "%Y-%m-%d")
    else:
        mes_ref = dt.datetime(2026, 9, 1) if a.exemplo else dt.datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    wb = Workbook()
    aba_leiame(wb)
    aba_parametros(wb, mes_ref)
    aba_clientes(wb)
    aba_pedidos(wb)
    aba_acompanhamento(wb)
    aba_atividades(wb)
    aba_custos(wb)
    aba_embalagens(wb)
    aba_uso(wb)
    aba_listas(wb)
    aba_painel(wb)
    if a.exemplo:
        dados_exemplo(wb)
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    print(saida)


if __name__ == "__main__":
    main()
