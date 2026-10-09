"""Monta um relatório uma vez e exporta em Word e Markdown; utilitários de planilha."""
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = "1F3864"
ENTRADA = "FFF2CC"   # células que o gerente preenche
CALC = "E2EFDA"      # fórmulas
ALERTA = "F8CBAD"
_fino = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


class Doc:
    def __init__(self):
        self.itens = []

    def h(self, txt, nivel=1):
        self.itens.append(("h", nivel, txt))

    def p(self, txt):
        self.itens.append(("p", txt))

    def lista(self, linhas):
        if linhas:
            self.itens.append(("l", list(linhas)))

    def tabela(self, df):
        if df is not None and len(df):
            self.itens.append(("t", df.fillna("").astype(str)))

    def md(self):
        out = []
        for it in self.itens:
            if it[0] == "h":
                out.append("#" * it[1] + " " + it[2])
            elif it[0] == "p":
                out.append(it[1])
            elif it[0] == "l":
                out.append("\n".join("- " + x for x in it[1]))
            else:
                df = it[1]
                t = "| " + " | ".join(df.columns) + " |\n|" + "---|" * len(df.columns)
                for row in df.itertuples(index=False):
                    t += "\n| " + " | ".join(str(v).replace("|", "/").replace("\n", " ") for v in row) + " |"
                out.append(t)
            out.append("")
        return "\n".join(out)

    def docx(self, caminho):
        d = Document()
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.5)
        d.styles["Normal"].font.name = "Calibri"
        d.styles["Normal"].font.size = Pt(10.5)
        for it in self.itens:
            if it[0] == "h":
                hd = d.add_heading(it[2], level=it[1])
                for r in hd.runs:
                    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            elif it[0] == "p":
                d.add_paragraph(it[1])
            elif it[0] == "l":
                for x in it[1]:
                    d.add_paragraph(x, style="List Bullet")
            else:
                df = it[1]
                t = d.add_table(rows=1, cols=len(df.columns))
                t.style = "Light Grid Accent 1"
                for j, c in enumerate(df.columns):
                    t.rows[0].cells[j].text = str(c)
                for row in df.itertuples(index=False):
                    cells = t.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = str(v)
                for row in t.rows:
                    for c in row.cells:
                        for par in c.paragraphs:
                            for r in par.runs:
                                r.font.size = Pt(9)
                d.add_paragraph()
        d.save(caminho)


def n(v, casas=1):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def rs(v):
    return "R$ " + n(v, 2)


def pct(v):
    return "-" if v is None or pd.isna(v) else f"{v * 100:.0f}%"


def dt(v):
    return "" if v is None or pd.isna(v) else pd.Timestamp(v).strftime("%d/%m/%Y")


# ------------------------------------------------------------------ planilha

def cabecalho(ws, colunas, linha=1, larguras=None, tipos=None):
    """tipos: lista com 'e' (entrada), 'c' (cálculo) ou None por coluna."""
    for j, nome in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=j, value=nome)
        tipo = (tipos or [None] * len(colunas))[j - 1]
        cor = {"e": "BF8F00", "c": "548235"}.get(tipo, AZUL)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=cor)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDA
        w = (larguras or {}).get(nome) or max(10, min(40, len(str(nome)) + 3))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[linha].height = 36
    ws.freeze_panes = ws.cell(row=linha + 1, column=2)


def pinta(ws, faixa, cor):
    for row in ws[faixa]:
        for c in row:
            c.fill = PatternFill("solid", fgColor=cor)
            c.border = BORDA


def lista_suspensa(ws, faixa, itens):
    dv = DataValidation(type="list", formula1='"' + ",".join(itens) + '"', allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(faixa)


def titulo(ws, txt, sub=None):
    ws["A1"] = txt
    ws["A1"].font = Font(bold=True, size=13, color=AZUL)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")
