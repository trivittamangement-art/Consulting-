"""Fechamento mensal do PPR Comex a partir das planilhas originais do gerente.

Uso:
    python3 fechar_mes.py --mes 2026-09 [--entrada ppr-comex/dados/origem] [--mestre ppr-comex/dados/PPR_Comex_2026.xlsx]

Lê os arquivos da pasta de entrada (periodicidade, S.A., controle de processos,
faturado x câmbio, uso de embalagens obsoletas, estoque de embalagens por cliente),
junta com o que o gerente já preencheu na planilha-mestre (realizado 2026, motivos,
conclusão das S.A., prazos, ações de custo) e gera:

    dados/PPR_Comex_<ano>.xlsx                 planilha-mestre atualizada (preserva o que foi digitado)
    relatorios/AAAA-MM/PPR_Comex_AAAA-MM.docx  relatório das metas do mês
    relatorios/AAAA-MM/PPR_Comex_AAAA-MM.md    o mesmo em texto
    relatorios/AAAA-MM/Resumo_gestor.txt       5 linhas para e-mail/WhatsApp
    relatorios/AAAA-MM/Pendencias.md           o que falta o gerente informar
"""
import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import ler_arquivos as L
from documento import (ALERTA, AZUL, CALC, ENTRADA, Doc, cabecalho, dt, lista_suspensa, n, pct, pinta,
                       rs, titulo)

RAIZ = Path(__file__).resolve().parent.parent
MESES_TXT = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
MESES_EXT = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro",
             "outubro", "novembro", "dezembro"]
MOTIVOS = ["Estoque alto no cliente", "Preço / câmbio", "Perdeu para concorrente", "Reclamação de qualidade",
           "Problema de embalagem", "Atraso em embarque anterior", "Barreira sanitária / regulatória",
           "Frete / contêiner / logística", "Demanda sazonal", "Crédito / financeiro do cliente",
           "Troca de comprador / estrutura", "Produto indisponível", "Pedido adiado para o mês seguinte",
           "Cliente encerrou atividade", "Sem retorno do cliente", "Não cumpriu a periodicidade", "Outro"]
TIPOS_ADMIN = ("COMPUTADOR", "IMPRESSORA", "SOFTWARE", "PERMISS", "SUPORTE", "TRANSPORTE DE COLAB", "MOTORISTA")
PRAZO_SUGERIDO = [("ARTE", 3), ("EMBALAGE", 7), ("AMOSTRA", 5), ("CADASTRO", 3), ("CANCELAMENTO", 1),
                  ("ALTERA", 1), ("DOCUMENTO", 5), ("MOTORISTA", 2), ("TRANSPORTE", 2)]
TIPOS_GANHO = ["Economia direta", "Custo evitado", "Custo transferido ao cliente", "Receita recuperada"]


# ------------------------------------------------------------------ arquivos de entrada

def localizar(pasta):
    achados = {"periodicidade": [], "sa": [], "cambio": [], "processos": [], "uso": [], "estoque": []}
    for f in sorted(Path(pasta).glob("*.xlsx")):
        nome = L.norm(re.sub(r"[_\-]+", " ", f.stem))
        if "periodicidade" in nome:
            achados["periodicidade"].append(f)
        elif "solicita" in nome and "atividade" in nome:
            achados["sa"].append(f)
        elif "faturado" in nome:
            achados["cambio"].append(f)
        elif "processos" in nome or "custos e vendas" in nome:
            achados["processos"].append(f)
        elif "uso de embalage" in nome:
            achados["uso"].append(f)
        elif "estoque" in nome:
            achados["estoque"].append(f)
    return achados


def apelidos(caminho):
    """dados/apelidos.csv (apelido;cliente) para casar nomes curtos com o cadastro."""
    if caminho.exists():
        df = pd.read_csv(caminho, sep=";", dtype=str).dropna()
        return {L.norm(a): c for a, c in zip(df.iloc[:, 0], df.iloc[:, 1])}
    caminho.write_text("apelido;cliente\n", encoding="utf-8")
    return {}


def preferencias(caminho):
    """dados/preferencias.json: decisões do gerente que valem para todos os meses."""
    padrao = {"sa_sem_data_conclusao": False, "descritivo_embalagens": True}
    if caminho.exists():
        padrao.update(json.loads(caminho.read_text(encoding="utf-8")))
    return padrao


def apelido(nome, ap):
    k = L.norm(nome)
    if k in ap:
        return ap[k]
    for a, c in ap.items():
        if k.startswith(a + " ") or k == a:
            return c
    return None


def eh_novo(nome, ap):
    """Cliente que o gerente confirmou como novo (apelidos.csv com cliente = NOVO)."""
    return str(apelido(nome, ap) or "").upper() == "NOVO"


def casar(nome, base, ap):
    """Nome do cadastro de periodicidade que corresponde a 'nome' (ou None)."""
    c = apelido(nome, ap)
    if c:
        return None if c.upper() == "NOVO" else c
    k = L.norm(nome)
    tok = re.split(r"[\s(]+", k)[0]
    cands = [b for b in base if re.split(r"[\s(]+", L.norm(b))[0] == tok]
    return cands[0] if len(cands) == 1 else None


# ------------------------------------------------------------------ planilha-mestre: ler o que o gerente digitou

def ler_mestre(caminho):
    m = {"realizado": {}, "evid_manual": {}, "motivos": pd.DataFrame(), "sa": {}, "prazos": {}, "acoes": pd.DataFrame()}
    if not caminho.exists():
        return m
    wb = load_workbook(caminho, data_only=False)
    if "M2 Realizado 2026" in wb.sheetnames or any(s.startswith("M2 Realizado") for s in wb.sheetnames):
        ws = wb[next(s for s in wb.sheetnames if s.startswith("M2 Realizado"))]
        for r in ws.iter_rows(min_row=4, values_only=True):
            if not r[0] or str(r[0]).startswith("Total"):
                continue
            for j in range(12):
                v = r[1 + j]
                if isinstance(v, (int, float)):
                    m["realizado"][(str(r[0]).strip(), j + 1)] = v
    if "M2 Motivos" in wb.sheetnames:
        m["motivos"] = pd.DataFrame(wb["M2 Motivos"].values)
        m["motivos"].columns = m["motivos"].iloc[0]
        m["motivos"] = m["motivos"].iloc[1:].dropna(how="all")
    if "M3a SA" in wb.sheetnames:
        ws = wb["M3a SA"]
        cab = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            d = dict(zip(cab, r))
            m["sa"][(str(d["Código"]), str(d["Descrição"])[:60])] = d
    if "M3a Prazos" in wb.sheetnames:
        for r in wb["M3a Prazos"].iter_rows(min_row=2, values_only=True):
            if r[0]:
                m["prazos"][r[0]] = (r[1], r[2])
    if "M3b Ações" in wb.sheetnames:
        a = pd.DataFrame(wb["M3b Ações"].values)
        a.columns = a.iloc[0]
        m["acoes"] = a.iloc[1:].dropna(how="all")
    return m


# ------------------------------------------------------------------ principal

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mes", required=True, help="AAAA-MM")
    ap.add_argument("--entrada", default=str(RAIZ / "dados" / "origem"))
    ap.add_argument("--mestre")
    ap.add_argument("--saida")
    a = ap.parse_args()

    ano, mes = map(int, a.mes.split("-"))
    base_ano = ano - 1
    nome_mes = f"{MESES_EXT[mes - 1]}/{ano}"
    mestre_path = Path(a.mestre) if a.mestre else RAIZ / "dados" / f"PPR_Comex_{ano}.xlsx"
    saida = Path(a.saida) if a.saida else RAIZ / "relatorios" / a.mes
    saida.mkdir(parents=True, exist_ok=True)
    arq = localizar(a.entrada)
    ap_map = apelidos(RAIZ / "dados" / "apelidos.csv")
    mest = ler_mestre(mestre_path)
    pref = preferencias(RAIZ / "dados" / "preferencias.json")
    pend, alertas, notas = [], [], []  # pendências do gerente; problemas nos dados; avisos sem ação

    # ================================================================ meta 2: positivação
    peds, cads = [], []
    for f in arq["periodicidade"]:
        p, c, _ = L.ler_periodicidade(f)
        p["arquivo"], c["arquivo"] = f.name, f.name
        peds.append(p)
        cads.append(c)
    ped = pd.concat(peds) if peds else pd.DataFrame(columns=["cliente", "ano", "mes", "qtd", "arquivo"])
    cad = pd.concat(cads) if cads else pd.DataFrame(columns=["cliente", "ano", "arquivo"])
    # para cada ano, fica o arquivo com mais pedidos lançados (o mais completo)
    melhor = ped.groupby(["ano", "arquivo"]).qtd.sum().reset_index().sort_values("qtd").drop_duplicates("ano", keep="last")
    escolha = dict(zip(melhor.ano, melhor.arquivo))
    ped = ped[ped.apply(lambda r: escolha.get(r.ano) == r.arquivo, axis=1)]
    cad = cad[cad.apply(lambda r: escolha.get(r.ano) == r.arquivo, axis=1)]
    base = ped[ped.ano == base_ano].pivot_table(index="cliente", columns="mes", values="qtd", aggfunc="sum").reindex(
        columns=range(1, 13), fill_value=0).fillna(0)
    cad_b = cad[cad.ano == base_ano].drop_duplicates("cliente").set_index("cliente")
    clientes_base = list(base.index)
    if not len(base):
        pend.append(f"Meta 2: enviar a planilha de periodicidade com os pedidos de {base_ano} por cliente e mês.")

    # evidências de embarque no ano (invoices nos arquivos de embalagem)
    uso = pd.concat([L.ler_uso_embalagens(f, ano) for f in arq["uso"]]) if arq["uso"] else pd.DataFrame()
    evid = {}
    novos = set()
    if len(uso):
        for (cli, m_), g in uso.groupby(["cliente", "mes"]):
            alvo = casar(cli, clientes_base, ap_map)
            invs = sorted({re.sub(r"\s+", "/", str(i)) for i in g.invoice.dropna()})
            chave = alvo or cli.strip()
            if not alvo:
                novos.add(chave)
            evid.setdefault(chave, {}).setdefault(m_, set()).update(invs)

    # realizado do ano: planilha de periodicidade do ano (se existir) + o que foi digitado na mestre
    real = ped[ped.ano == ano].pivot_table(index="cliente", columns="mes", values="qtd", aggfunc="sum") \
        if (ped.ano == ano).any() else pd.DataFrame()
    realizado = {}
    for cli in real.index:
        for m_ in real.columns:
            if real.loc[cli, m_]:
                realizado[(cli, int(m_))] = real.loc[cli, m_]
    realizado.update(mest["realizado"])
    tem_realizado = any(k[1] <= mes for k in realizado)
    clientes_ano = clientes_base + sorted(novos | {k[0] for k in realizado if k[0] not in clientes_base})

    meta_mes = (base.sum() * 1.2).round(1)
    real_mes = pd.Series({m_: sum(v for (c, mm), v in realizado.items() if mm == m_) for m_ in range(1, 13)})
    pos_2025 = int((base.sum(axis=1) > 0).sum())
    sem_compra = [c for c in clientes_base if base.loc[c].sum() == 0]
    top = base.sum(axis=1).sort_values(ascending=False)
    concentr = top.iloc[0] / top.sum() if top.sum() else 0
    esperados = [c for c in clientes_base if base.loc[c, mes] > 0]
    prox = {m_: [c for c in clientes_base if base.loc[c, m_] > 0] for m_ in range(mes + 1, min(mes + 3, 12) + 1)}

    if not tem_realizado:
        pend.append(f"Meta 2: lançar os pedidos de {ano} por cliente e mês (aba 'M2 Realizado {ano}' da planilha-mestre, "
                    f"ou mandar a planilha de periodicidade de {ano}). Sem isso não dá para medir a meta de positivação.")
    nao_base = sorted(novos)
    novos_conf = [c for c in nao_base if eh_novo(c, ap_map)]
    novos_duv = [c for c in nao_base if c not in novos_conf]
    fat_fora, fat_novo = [], []
    for f in arq["cambio"]:
        for cli in sorted(L.ler_faturado_cambio(f).cliente.unique()):
            if eh_novo(cli, ap_map):
                fat_novo.append(cli)
            elif not casar(cli, clientes_base, ap_map):
                fat_fora.append(cli)
    if fat_fora:
        alertas.append(f"Clientes faturados em {base_ano} (Faturado x Câmbio) que não aparecem na periodicidade de {base_ano}: "
                       + ", ".join(fat_fora) + ". Se compraram em " + str(base_ano) + ", a base de pedidos está incompleta.")
    if fat_novo:
        alertas.append(f"Clientes confirmados como novos na positivação, mas com faturamento em {base_ano} no Faturado x "
                       f"Câmbio: {', '.join(fat_novo)}. Ter a justificativa pronta caso a apuração do PPR cruze com o "
                       "faturamento.")
    if novos_duv:
        alertas.append(f"Clientes com embarque em {ano} que não estão na planilha de periodicidade de {base_ano}: "
                       + ", ".join(novos_duv) + ". Se forem clientes novos, contam a favor dos +20%; se forem nomes "
                       "diferentes de um cliente já cadastrado, informe em dados/apelidos.csv.")

    # ================================================================ meta 3a: S.A.
    sa = pd.concat([L.ler_sa(f)[0].assign(arquivo=f.name) for f in arq["sa"]]) if arq["sa"] else pd.DataFrame()
    usuarios = pd.concat([L.ler_sa(f)[1] for f in arq["sa"]]) if arq["sa"] else pd.DataFrame()
    if len(sa):
        sa["natureza"] = sa["tipo"].map(lambda t: "Apoio administrativo" if any(k in t for k in TIPOS_ADMIN)
                                        else "Processo de exportação")
        sa["pedidos_citados"] = sa["descricao"].map(lambda d: ", ".join(sorted(set(re.findall(r"\b2\d{5}\b", d)))))
        for t in sa["tipo"].unique():
            if t not in mest["prazos"]:
                pz = next((d for k, d in PRAZO_SUGERIDO if k in t), 3)
                mest["prazos"][t] = (pz, "Sugestão: confirmar com o setor")
        sa["prazo"] = sa["tipo"].map(lambda t: mest["prazos"][t][0])
        # o que o gerente já digitou na mestre
        for campo, col_m in (("conc_m", "Data conclusão"), ("cobrancas", "Nº cobranças"),
                             ("ultima_cobranca", "Última cobrança"), ("obs", "Observação"),
                             ("pedido_vinc", "Pedido/Invoice vinculado"), ("nat_m", "Natureza")):
            sa[campo] = [mest["sa"].get((str(r.codigo), str(r.descricao)[:60]), {}).get(col_m) for r in sa.itertuples()]
        sa["data_conclusao"] = pd.to_datetime(sa["data_conclusao"], errors="coerce").fillna(
            pd.to_datetime(sa["conc_m"], errors="coerce"))
        sa.loc[sa["nat_m"].notna(), "natureza"] = sa["nat_m"]
        sa_mes = sa[(sa.data_criacao.dt.year == ano) & (sa.data_criacao.dt.month == mes)]
        unicas = sa_mes[~sa_mes.duplicada].drop_duplicates("codigo")
        if sa_mes["duplicada"].any():
            alertas.append(f"S.A. repetida na exportação: {', '.join(map(str, sa_mes[sa_mes.duplicada].codigo))} "
                           "(contada uma vez só).")
        if sa_mes["mesmo_codigo_outra_aba"].any():
            cods = sorted(set(sa_mes[sa_mes.mesmo_codigo_outra_aba].codigo))
            alertas.append(f"Mesmo código de S.A. em duas abas/setores com descrições diferentes: {', '.join(map(str, cods))}.")
        titulos = {str(t) for f in arq["sa"] for t in [load_workbook(f, read_only=True)[w].cell(1, 1).value
                                                        for w in load_workbook(f, read_only=True).sheetnames] if t}
        if any(MESES_EXT[mes - 1][:3] not in L.norm(t) for t in titulos if "relatorio" in L.norm(t)):
            alertas.append("Abas da exportação de S.A. com título de outro mês (ex.: 'RELATÓRIO DE SA - JUNHO' "
                           "em dados de setembro). Ajustar o modelo do relatório.")
        sem_conclusao = unicas["data_conclusao"].isna().all()
        if sem_conclusao and pref["sa_sem_data_conclusao"]:
            notas.append("O sistema não exporta a data de conclusão das S.A. O prazo por setor passa a ser medido pela "
                         "aba 'M3a SA': anotar a data de entrega e as cobranças das S.A. de processo de exportação.")
        elif sem_conclusao:
            pend.append("Meta 3a: a exportação de S.A. não traz status nem data de conclusão. Exportar do sistema "
                        "com as colunas 'Status' e 'Data de conclusão' (ou preencher na aba 'M3a SA'), e anotar as "
                        "cobranças feitas (quantas e quando). Sem isso não dá para medir prazo por setor.")
    else:
        sa_mes = unicas = pd.DataFrame()
        pend.append("Meta 3a: enviar a exportação de S.A. do mês.")

    fluxo = pd.concat([L.ler_fluxo(f) for f in arq["processos"]]) if arq["processos"] else pd.DataFrame()
    if len(fluxo):
        if fluxo["data_pedido_suspeita"].any():
            alertas.append("Controle de Fluxo: data do pedido posterior à saída do navio/previsão em "
                           f"{int(fluxo['data_pedido_suspeita'].sum())} processo(s) (provável erro de digitação do ano); "
                           "fórmulas da planilha original geram números como -45.694 quando falta data.")

    # ================================================================ meta 3b: custos
    gastos = pd.concat([L.ler_gastos(f) for f in arq["processos"]]) if arq["processos"] else pd.DataFrame()
    cambio = pd.concat([L.ler_faturado_cambio(f) for f in arq["cambio"]]) if arq["cambio"] else pd.DataFrame()
    if len(cambio) and cambio["duplicada"].any():
        d = cambio[cambio.duplicada]
        alertas.append("Faturado x Câmbio: invoice lançada duas vezes (" + ", ".join(f"{r.invoice} {r.cliente}"
                       for r in d.itertuples()) + f"); o total da planilha fica inflado em {rs(d.recebido_cambio.sum())}.")

    # ================================================================ meta 3c: embalagens
    estoque_parado = {}
    for f in arq["uso"]:
        estoque_parado = L.ler_estoque_parado(f) or estoque_parado
    est_cli = pd.concat([L.ler_estoque_cliente(f, ano).assign(cliente=re.sub(r"^.*?Estoque[_ ]Embalagens[_ ]?", "",
                         f.stem).replace("_", " ")) for f in arq["estoque"]]) if arq["estoque"] else pd.DataFrame()
    uso_mes = uso[uso.mes == mes] if len(uso) else pd.DataFrame()
    if len(uso):
        susp = uso[(uso.qtd_embalagem.notna()) & (uso.codigo_embalagem.astype(str) == uso.qtd_embalagem.map(
            lambda v: str(int(v)) if pd.notna(v) and float(v).is_integer() else ""))]
        if len(susp):
            alertas.append("Uso de embalagens: quantidade igual ao código da embalagem em " + ", ".join(
                f"{MESES_TXT[r.mes - 1]} {r.cliente} cód. {r.codigo_embalagem}" for r in susp.itertuples())
                + " (provável erro de digitação).")
        if uso["celula_com_varios_valores"].any():
            alertas.append("Uso de embalagens: várias células com dois ou três códigos/valores juntos "
                           "('22469 / 22468', 'R$ 430,83 R$ 114'). A planilha-mestre separa em uma linha por código.")
    if len(uso_mes) == 0:
        pend.append(f"Meta 3c: enviar a aba de {MESES_EXT[mes - 1]} do 'Uso de Embalagens Obsoletas'.")
    vp = estoque_parado.get("valor", pd.DataFrame())
    ultimo_aging = vp["data"].max() if len(vp) else None
    if ultimo_aging is not None and (ultimo_aging.year, ultimo_aging.month) < (ano, mes) and not pref["descritivo_embalagens"]:
        notas.append(f"Estoque de embalagens por faixa de dias: último retrato de {dt(ultimo_aging)[3:]}.")
    elif ultimo_aging is not None and (ultimo_aging.year, ultimo_aging.month) < (ano, mes):
        pend.append(f"Meta 3c: atualizar o estoque de embalagens por faixa de dias sem movimentação; o último é de "
                    f"{dt(ultimo_aging)[3:]}. Melhor ainda: o relatório por código (código, descrição, quantidade, valor, "
                    "dias sem movimentação), para casar embalagem parada com pedido em carteira.")
    if len(est_cli):
        piv = est_cli.groupby(["codigo", "mes"]).estoque.sum().unstack()
        salto = piv[(piv.max(axis=1) > 0) & (piv.max(axis=1) / piv.replace(0, pd.NA).min(axis=1) > 100)]
        if len(salto):
            alertas.append("Estoque por cliente: códigos com saltos de mais de 100 vezes entre meses (" +
                           ", ".join(map(str, salto.index)) + "), sinal de unidade misturada (kg, milheiro, unidade).")

    # ================================================================ planilha-mestre
    wb = Workbook()
    ws = wb.active
    ws.title = "Painel"

    # --- M2 Base
    wsb = wb.create_sheet(f"M2 Base {base_ano}")
    titulo(wsb, f"Positivações (pedidos) de {base_ano} por cliente e mês", f"Fonte: {escolha.get(base_ano, '-')}")
    cols = ["Cliente"] + MESES_TXT + ["Total", "Frequência", "Periodicidade (meses)", f"Cumprido {base_ano}?", "Observação"]
    cabecalho(wsb, cols, linha=3, larguras={"Cliente": 40, "Observação": 50})
    for i, c in enumerate(clientes_base, start=4):
        wsb.cell(row=i, column=1, value=c)
        for m_ in range(1, 13):
            wsb.cell(row=i, column=1 + m_, value=float(base.loc[c, m_]))
        wsb.cell(row=i, column=14, value=f"=SUM(B{i}:M{i})")
        info = cad_b.loc[c] if c in cad_b.index else {}
        wsb.cell(row=i, column=15, value=info.get("frequencia_txt") if len(info) else None)
        wsb.cell(row=i, column=16, value=info.get("periodicidade_meses") if len(info) else None)
        wsb.cell(row=i, column=17, value=info.get("cumprido") if len(info) else None)
        wsb.cell(row=i, column=18, value=info.get("observacao") if len(info) else None)
    ult_b = 3 + len(clientes_base)
    tb = ult_b + 1
    wsb.cell(row=tb, column=1, value=f"Total {base_ano}").font = Font(bold=True)
    wsb.cell(row=tb + 1, column=1, value=f"Meta {ano} (+20%)").font = Font(bold=True)
    wsb.cell(row=tb + 2, column=1, value="Clientes positivados").font = Font(bold=True)
    for j in range(2, 15):
        L_ = get_column_letter(j)
        wsb.cell(row=tb, column=j, value=f"=SUM({L_}4:{L_}{ult_b})")
        wsb.cell(row=tb + 1, column=j, value=f"=ROUND({L_}{tb}*1.2,1)")
        wsb.cell(row=tb + 2, column=j, value=f'=COUNTIF({L_}4:{L_}{ult_b},">0")')
    pinta(wsb, f"A{tb}:N{tb + 2}", CALC)

    # --- M2 Realizado
    wsr = wb.create_sheet(f"M2 Realizado {ano}")
    titulo(wsr, f"Positivações (pedidos) de {ano} por cliente e mês: PREENCHER as células amarelas",
           "Coluna 'Evidência' = invoices encontradas nos arquivos de embalagem; serve de ajuda, não substitui o lançamento.")
    cols = ["Cliente"] + MESES_TXT + ["Total", "Na base?", "Evidência nos arquivos (mês: invoices)"]
    cabecalho(wsr, cols, linha=3, larguras={"Cliente": 40, "Evidência nos arquivos (mês: invoices)": 60},
              tipos=[None] + ["e"] * 12 + ["c", None, None])
    for i, c in enumerate(clientes_ano, start=4):
        wsr.cell(row=i, column=1, value=c)
        for m_ in range(1, 13):
            v = realizado.get((c, m_))
            cel = wsr.cell(row=i, column=1 + m_, value=v)
            cel.fill = PatternFill("solid", fgColor=ENTRADA)
        wsr.cell(row=i, column=14, value=f"=SUM(B{i}:M{i})").fill = PatternFill("solid", fgColor=CALC)
        wsr.cell(row=i, column=15, value="Sim" if c in clientes_base else "Não (novo)")
        ev = evid.get(c, {})
        wsr.cell(row=i, column=16, value="; ".join(f"{MESES_TXT[k - 1]}: {', '.join(sorted(v))}" for k, v in sorted(ev.items())))
    ult_r = 3 + len(clientes_ano)
    tr = ult_r + 1
    wsr.cell(row=tr, column=1, value=f"Total {ano}").font = Font(bold=True)
    wsr.cell(row=tr + 1, column=1, value="Clientes positivados").font = Font(bold=True)
    for j in range(2, 15):
        L_ = get_column_letter(j)
        wsr.cell(row=tr, column=j, value=f"=SUM({L_}4:{L_}{ult_r})")
        wsr.cell(row=tr + 1, column=j, value=f'=COUNTIF({L_}4:{L_}{ult_r},">0")')
    pinta(wsr, f"A{tr}:N{tr + 1}", CALC)

    # --- M2 Meta x Realizado
    wsm = wb.create_sheet("M2 Meta x Realizado")
    titulo(wsm, f"Meta de positivação {ano}: base {base_ano} + 20%, mês a mês e acumulado")
    cols = ["Mês", f"Pedidos {base_ano}", f"Meta {ano}", f"Realizado {ano}", "% do mês", f"Meta acumulada",
            "Realizado acumulado", "% acumulado", f"Clientes {base_ano}", f"Clientes {ano}"]
    cabecalho(wsm, cols, linha=3, tipos=[None] + ["c"] * 9)
    B, R = f"'M2 Base {base_ano}'", f"'M2 Realizado {ano}'"
    for m_ in range(1, 13):
        i = 3 + m_
        L_ = get_column_letter(1 + m_)
        wsm.cell(row=i, column=1, value=MESES_TXT[m_ - 1])
        wsm.cell(row=i, column=2, value=f"={B}!{L_}{tb}")
        wsm.cell(row=i, column=3, value=f"={B}!{L_}{tb + 1}")
        wsm.cell(row=i, column=4, value=f"={R}!{L_}{tr}")
        wsm.cell(row=i, column=5, value=f'=IF(C{i}=0,"",D{i}/C{i})').number_format = "0%"
        wsm.cell(row=i, column=6, value=f"=SUM($C$4:C{i})")
        wsm.cell(row=i, column=7, value=f"=SUM($D$4:D{i})")
        wsm.cell(row=i, column=8, value=f'=IF(F{i}=0,"",G{i}/F{i})').number_format = "0%"
        wsm.cell(row=i, column=9, value=f"={B}!{L_}{tb + 2}")
        wsm.cell(row=i, column=10, value=f"={R}!{L_}{tr + 1}")
    wsm.cell(row=16, column=1, value="Ano").font = Font(bold=True)
    for j, f_ in ((2, "=SUM(B4:B15)"), (3, "=SUM(C4:C15)"), (4, "=SUM(D4:D15)"), (5, '=IF(C16=0,"",D16/C16)')):
        wsm.cell(row=16, column=j, value=f_)
    wsm["E16"].number_format = "0%"

    # --- M2 Acompanhamento do mês
    wsa = wb.create_sheet("M2 Acompanhamento")
    titulo(wsa, "Situação de cada cliente no mês de referência (mês em Painel!B3)",
           "Previsto = comprou no mesmo mês em " + str(base_ano) + ". Motivo, ação e retorno vão para a aba 'M2 Motivos'.")
    cols = ["Cliente", "Frequência", f"Meses de compra em {base_ano}", f"Pedidos {base_ano} no mês",
            f"Pedidos {ano} no mês", "Situação no mês", f"Pedidos {ano} no ano", "Próximos meses previstos"]
    cabecalho(wsa, cols, linha=3, larguras={"Cliente": 40, f"Meses de compra em {base_ano}": 30,
                                             "Próximos meses previstos": 26}, tipos=[None, None, None, "c", "c", "c", "c", None])
    for i, c in enumerate(clientes_ano, start=4):
        wsa.cell(row=i, column=1, value=c)
        if c in clientes_base:
            info = cad_b.loc[c] if c in cad_b.index else {}
            wsa.cell(row=i, column=2, value=info.get("frequencia_txt") if len(info) else None)
            wsa.cell(row=i, column=3, value=", ".join(MESES_TXT[m_ - 1] for m_ in range(1, 13) if base.loc[c, m_] > 0) or "não comprou")
            rb = clientes_base.index(c) + 4
            wsa.cell(row=i, column=4, value=f"=INDEX({B}!B{rb}:M{rb},Painel!$B$3)")
        else:
            wsa.cell(row=i, column=2, value="novo")
            wsa.cell(row=i, column=4, value=0)
        rr = clientes_ano.index(c) + 4
        wsa.cell(row=i, column=5, value=f"=INDEX({R}!B{rr}:M{rr},Painel!$B$3)")
        wsa.cell(row=i, column=6, value=(f'=IF(ISBLANK(INDEX({R}!B{rr}:M{rr},Painel!$B$3)),'
                                         f'IF(D{i}>0,"Previsto (sem lançamento)","Sem previsão"),'
                                         f'IF(N(E{i})>0,IF(D{i}>0,"Recomprou","Positivação extra"),'
                                         f'IF(D{i}>0,"Não recomprou","Sem previsão")))'))
        wsa.cell(row=i, column=7, value=f"={R}!N{rr}")
        if c in clientes_base:
            futuros = [MESES_TXT[m_ - 1] for m_ in range(mes + 1, 13) if base.loc[c, m_] > 0]
            wsa.cell(row=i, column=8, value=", ".join(futuros))
    pinta(wsa, f"D4:G{3 + len(clientes_ano)}", CALC)

    # --- M2 Motivos (histórico, editável)
    wsx = wb.create_sheet("M2 Motivos")
    cols = ["Mês", "Cliente", "Situação", "Motivo", "Detalhe", "Ação combinada", "Data de retorno", "Chance de recompra",
            "Status"]
    cabecalho(wsx, cols, larguras={"Cliente": 40, "Detalhe": 50, "Ação combinada": 40, "Motivo": 30},
              tipos=[None, None, None, "e", "e", "e", "e", "e", "e"])
    mot = mest["motivos"]
    linhas_mot = [list(r) for r in mot.itertuples(index=False)] if len(mot) else []
    ja = {(str(r[0]), str(r[1])) for r in linhas_mot}
    rot = f"{ano}-{mes:02d}"
    for c in esperados:
        if (rot, c) not in ja and realizado.get((c, mes), 0) == 0:
            linhas_mot.append([rot, c, "Não recomprou" if tem_realizado else "Previsto (aguardando lançamento)",
                               None, None, None, None, None, "Aberta"])
    for i, r in enumerate(linhas_mot, start=2):
        for j, v in enumerate(r[:9], start=1):
            wsx.cell(row=i, column=j, value=v)
    pinta(wsx, f"D2:I{max(2, len(linhas_mot) + 1)}", ENTRADA)
    lista_suspensa(wsx, "D2:D500", MOTIVOS)
    lista_suspensa(wsx, "H2:H500", ["Alta", "Média", "Baixa"])
    lista_suspensa(wsx, "I2:I500", ["Aberta", "Concluída", "Cancelada"])

    # --- M3a SA
    wss = wb.create_sheet("M3a SA")
    cols = ["Código", "Data criação", "Setor", "Tipo", "Descrição", "Usuário", "Natureza", "Pedido/Invoice vinculado",
            "Prazo (dias)", "Data limite", "Data conclusão", "Status", "Dias de atraso", "Nº cobranças",
            "Última cobrança", "Observação"]
    tipos = [None, None, None, None, None, None, "e", "e", "c", "c", "e", "c", "c", "e", "e", "e"]
    cabecalho(wss, cols, larguras={"Descrição": 60, "Tipo": 30, "Setor": 26}, tipos=tipos)
    if len(sa):
        sa_out = sa[~sa.duplicada].copy()
        outra = sa_out[sa_out.duplicated("codigo", keep="first")]
        sa_out = sa_out.drop_duplicates("codigo")
        for cod, aba_ in zip(outra.codigo, outra.aba):
            sa_out.loc[(sa_out.codigo == cod) & sa_out.obs.isna(), "obs"] = f"Mesmo código também na aba {aba_}"
    else:
        sa_out = sa
    for i, r in enumerate(sa_out.itertuples(), start=2):
        obs = r.obs
        if r.mesmo_codigo_outra_aba and not obs:
            obs = "Mesmo código aparece em outra aba"
        vals = [r.codigo, r.data_criacao.to_pydatetime() if pd.notna(r.data_criacao) else None, r.setor, r.tipo,
                r.descricao, r.usuario, r.natureza, r.pedido_vinc or r.pedidos_citados or None,
                f"=IFERROR(VLOOKUP(D{i},'M3a Prazos'!A:B,2,0),3)", f'=IF(B{i}="","",B{i}+I{i})',
                r.data_conclusao.to_pydatetime() if pd.notna(r.data_conclusao) else None,
                f'=IF(B{i}="","",IF(K{i}="",IF(TODAY()>J{i},"Vencida sem baixa","Em aberto"),IF(K{i}<=J{i},"No prazo","Com atraso")))',
                f'=IF(B{i}="","",MAX(0,IF(K{i}="",TODAY(),K{i})-J{i}))', r.cobrancas, r.ultima_cobranca, obs]
        for j, v in enumerate(vals, start=1):
            wss.cell(row=i, column=j, value=v)
        for j in (2, 10, 11, 15):
            wss.cell(row=i, column=j).number_format = "dd/mm/yyyy"
    nsa = len(sa_out) + 1
    for j, cor in ((7, ENTRADA), (8, ENTRADA), (11, ENTRADA), (14, ENTRADA), (15, ENTRADA), (16, ENTRADA),
                   (9, CALC), (10, CALC), (12, CALC), (13, CALC)):
        L_ = get_column_letter(j)
        pinta(wss, f"{L_}2:{L_}{max(2, nsa)}", cor)
    lista_suspensa(wss, "G2:G1000", ["Processo de exportação", "Apoio administrativo"])

    wsp = wb.create_sheet("M3a Prazos")
    cabecalho(wsp, ["Tipo de atividade", "Prazo (dias)", "Origem do prazo"], larguras={"Tipo de atividade": 45, "Origem do prazo": 34},
              tipos=[None, "e", None])
    for i, (t, (pz, origem)) in enumerate(sorted(mest["prazos"].items()), start=2):
        wsp.cell(row=i, column=1, value=t)
        wsp.cell(row=i, column=2, value=pz).fill = PatternFill("solid", fgColor=ENTRADA)
        wsp.cell(row=i, column=3, value=origem)

    wsr2 = wb.create_sheet("M3a Resumo setores")
    titulo(wsr2, "S.A. por setor (todas as S.A. da aba M3a SA)")
    cols = ["Setor", "S.A.", "De processo de exportação", "Concluídas", "No prazo", "% no prazo", "Vencidas sem baixa",
            "Cobranças"]
    cabecalho(wsr2, cols, linha=3, larguras={"Setor": 34}, tipos=[None] + ["c"] * 7)
    setores = sorted(sa_out["setor"].unique()) if len(sa_out) else []
    for i, s_ in enumerate(setores, start=4):
        wsr2.cell(row=i, column=1, value=s_)
        rg = "'M3a SA'!"
        wsr2.cell(row=i, column=2, value=f'=COUNTIF({rg}C:C,A{i})')
        wsr2.cell(row=i, column=3, value=f'=COUNTIFS({rg}C:C,A{i},{rg}G:G,"Processo de exportação")')
        wsr2.cell(row=i, column=4, value=f'=COUNTIFS({rg}C:C,A{i},{rg}K:K,"<>")')
        wsr2.cell(row=i, column=5, value=f'=COUNTIFS({rg}C:C,A{i},{rg}L:L,"No prazo")')
        wsr2.cell(row=i, column=6, value=f'=IF(D{i}=0,"",E{i}/D{i})').number_format = "0%"
        wsr2.cell(row=i, column=7, value=f'=COUNTIFS({rg}C:C,A{i},{rg}L:L,"Vencida sem baixa")')
        wsr2.cell(row=i, column=8, value=f'=SUMIF({rg}C:C,A{i},{rg}N:N)')

    wsf = wb.create_sheet("M3a Fluxo processos")
    titulo(wsf, "Controle de fluxo dos processos de exportação (tempos recalculados pelas datas)",
           "Em branco = falta uma das datas da etapa. 'Data suspeita' = data do pedido depois da saída do navio.")
    if len(fluxo):
        cols_f = ["invoice", "pedido", "cliente", "pais", "data_pedido", "previsao_saida", "saida_navio", "tinha_estoque",
                  "status", "motivo", "dias_total", "dias_vs_previsao", "data_pedido_suspeita"]
        etapas = [c for c in fluxo.columns if c not in cols_f]
        cab_f = ["Invoice", "Pedido", "Cliente", "País", "Data do pedido", "Previsão saída", "Saída do navio",
                 "Tinha estoque?", "Status", "Motivo do atraso", "Dias pedido → navio", "Dias vs previsão",
                 "Data suspeita"] + etapas
        cabecalho(wsf, cab_f, linha=4, larguras={"Cliente": 24, "Motivo do atraso": 36})
        for i, r in enumerate(fluxo[cols_f + etapas].itertuples(index=False), start=5):
            for j, v in enumerate(r, start=1):
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime() if pd.notna(v) else None
                elif v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
                    v = None
                elif isinstance(v, (bool,)) or type(v).__name__ == "bool_":
                    v = "Sim" if v else ""
                c_ = wsf.cell(row=i, column=j, value=v)
                if j in (5, 6, 7):
                    c_.number_format = "dd/mm/yyyy"

    # --- M3b
    wsc = wb.create_sheet("M3b Ações")
    cols = ["Mês", "Categoria", "Descrição da ação", "Tipo de ganho", "Custo de referência (R$)", "Custo realizado (R$)",
            "Economia (R$)", "Evidência", "Origem"]
    cabecalho(wsc, cols, larguras={"Descrição da ação": 60, "Evidência": 40, "Categoria": 26},
              tipos=["e", "e", "e", "e", "e", "e", "c", "e", None])
    acoes = []
    if len(mest["acoes"]):
        acoes = [list(r)[:9] for r in mest["acoes"].itertuples(index=False) if (list(r) + [None] * 9)[8] != "Automático"]
    if len(uso):
        for m_, g in uso.groupby("mes"):
            if m_ > mes:
                continue
            acoes.append([f"{ano}-{m_:02d}", "Embalagens reaproveitadas",
                          "Embalagens obsoletas usadas nos pedidos " + ", ".join(sorted({f"{r.cliente} {r.invoice}".replace("\n", "/")
                                                                                         for r in g.itertuples()})),
                          "Custo evitado", g.valor.sum(), 0, None, f"Uso de Embalagens Obsoletas - aba {MESES_TXT[m_ - 1]}",
                          "Automático"])
    for i, r in enumerate(acoes, start=2):
        for j, v in enumerate(r, start=1):
            if j == 7:
                v = f'=IF(AND(E{i}="",F{i}=""),"",N(E{i})-N(F{i}))'
            wsc.cell(row=i, column=j, value=v)
        for j in (5, 6, 7):
            wsc.cell(row=i, column=j).number_format = "#,##0.00"
    lista_suspensa(wsc, "D2:D500", TIPOS_GANHO)
    lista_suspensa(wsc, "B2:B500", ["Amostras", "Frete de amostras", "Inspeções", "Viagens internacionais",
                                    "Impressões / documentos", "Courier", "Frete internacional", "Comissão de agente",
                                    "Embalagens reaproveitadas", "Motorista / deslocamentos", "Outros"])
    pinta(wsc, f"A{len(acoes) + 2}:F{len(acoes) + 40}", ENTRADA)

    wsg = wb.create_sheet(f"M3b Custos {base_ano}")
    titulo(wsg, "Custos vinculados por processo (fonte: Controle de Gastos)")
    if len(gastos):
        cg = [c for c in gastos.columns]
        cabecalho(wsg, cg, linha=3)
        for i, r in enumerate(gastos.itertuples(index=False), start=4):
            for j, v in enumerate(r, start=1):
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime() if pd.notna(v) else None
                elif isinstance(v, float) and pd.isna(v):
                    v = None
                wsg.cell(row=i, column=j, value=v)
    wsk = wb.create_sheet(f"M3b Câmbio {base_ano}")
    titulo(wsk, "Faturado x recebido no câmbio (linhas repetidas marcadas e fora do total)")
    if len(cambio):
        cabecalho(wsk, ["Invoice", "Cliente", "Valor NF", "Data NF", "Recebido câmbio", "Variação", "Variação %",
                        "Repetida?", "Data digitada (texto)"], linha=3)
        for i, r in enumerate(cambio.itertuples(), start=4):
            vals = [r.invoice, r.cliente, r.valor_nf, r.data_nf.to_pydatetime() if pd.notna(r.data_nf) else None,
                    r.recebido_cambio, f"=E{i}-C{i}", f'=IF(C{i}=0,"",F{i}/C{i})', "Sim" if r.duplicada else "",
                    r.data_nf_texto]
            for j, v in enumerate(vals, start=1):
                wsk.cell(row=i, column=j, value=v)
            wsk.cell(row=i, column=7).number_format = "0.0%"
            if r.duplicada:
                pinta(wsk, f"A{i}:I{i}", ALERTA)
        f_ = 4 + len(cambio)
        wsk.cell(row=f_, column=1, value="Total sem repetidas").font = Font(bold=True)
        for j, L_ in ((3, "C"), (5, "E"), (6, "F")):
            wsk.cell(row=f_, column=j, value=f'=SUMIF($H$4:$H${f_ - 1},"<>Sim",{L_}4:{L_}{f_ - 1})')

    # --- M3c
    wsu = wb.create_sheet(f"M3c Uso {ano}")
    titulo(wsu, "Uso de embalagens obsoletas: uma linha por código (células com vários códigos já separadas)")
    cols = ["Mês", "Cliente", "Invoice", "Caixas do produto", "Descrição do produto", "Código embalagem", "Referência",
            "Tipo", "Qtd embalagem", "Valor (R$)", "Dias parada", "Origem: célula com vários valores"]
    cabecalho(wsu, cols, linha=3, larguras={"Descrição do produto": 46})
    if len(uso):
        for i, r in enumerate(uso.sort_values(["mes"]).itertuples(), start=4):
            vals = [MESES_TXT[r.mes - 1], r.cliente, str(r.invoice).replace("\n", "/"), r.qtd_caixas_produto, r.descricao,
                    r.codigo_embalagem, r.referencia, (r.tipo or "").strip(), r.qtd_embalagem, r.valor, r.dias_parada,
                    "Sim" if r.celula_com_varios_valores else ""]
            for j, v in enumerate(vals, start=1):
                wsu.cell(row=i, column=j, value=None if isinstance(v, float) and pd.isna(v) else v)
            wsu.cell(row=i, column=10).number_format = "#,##0.00"
    wsur = wb.create_sheet("M3c Resumo")
    titulo(wsur, "Embalagens obsoletas reaproveitadas por mês")
    cabecalho(wsur, ["Mês", "Valor reaproveitado (R$)", "Itens", "Acumulado no ano (R$)"], linha=3, tipos=[None, "c", "c", "c"])
    U = f"'M3c Uso {ano}'"
    for m_ in range(1, 13):
        i = 3 + m_
        wsur.cell(row=i, column=1, value=MESES_TXT[m_ - 1])
        wsur.cell(row=i, column=2, value=f'=SUMIF({U}!A:A,A{i},{U}!J:J)').number_format = "#,##0.00"
        wsur.cell(row=i, column=3, value=f'=COUNTIFS({U}!A:A,A{i},{U}!J:J,">0")')
        wsur.cell(row=i, column=4, value=f"=SUM($B$4:B{i})").number_format = "#,##0.00"
    if len(vp):
        wse = wb.create_sheet("M3c Estoque parado")
        titulo(wse, "Estoque total de embalagens por faixa de dias sem movimentação (R$)")
        pv = vp[vp.bloco.str.startswith("Descrição Estoque Total")].pivot_table(index="faixa", columns="data", values="valor",
                                                                                 aggfunc="sum", sort=False)
        cabecalho(wse, ["Faixa"] + [d.strftime("%m/%Y") for d in pv.columns], linha=3, larguras={"Faixa": 36})
        for i, (fx, row) in enumerate(pv.iterrows(), start=4):
            wse.cell(row=i, column=1, value=fx)
            for j, v in enumerate(row, start=2):
                wse.cell(row=i, column=j, value=v).number_format = "#,##0"
    if len(est_cli):
        wsc2 = wb.create_sheet("M3c Estoque por cliente")
        titulo(wsc2, "Estoque de embalagem dedicada por cliente, por mês")
        pv = est_cli.pivot_table(index=["cliente", "grupo", "codigo", "descricao"], columns="mes", values="estoque",
                                 aggfunc="sum")
        cabecalho(wsc2, ["Cliente", "Grupo", "Código", "Descrição"] + [MESES_TXT[m_ - 1] for m_ in pv.columns],
                  linha=3, larguras={"Descrição": 40, "Grupo": 34})
        for i, (k, row) in enumerate(pv.iterrows(), start=4):
            for j, v in enumerate(k, start=1):
                wsc2.cell(row=i, column=j, value=v)
            for j, v in enumerate(row, start=5):
                wsc2.cell(row=i, column=j, value=None if pd.isna(v) else v)

    # --- Pendências e Painel
    wpd = wb.create_sheet("Pendências")
    cabecalho(wpd, ["#", "O que falta", "Tipo"], larguras={"O que falta": 140})
    for i, t in enumerate(pend + alertas, start=2):
        wpd.cell(row=i, column=1, value=i - 1)
        wpd.cell(row=i, column=2, value=t)
        wpd.cell(row=i, column=3, value="Informar" if t in pend else "Corrigir na planilha de origem")

    titulo(ws, "Painel PPR Comex", "Troque o mês em B3 e todas as abas M2 se ajustam. Amarelo = preencher; verde = fórmula.")
    ws["A3"], ws["B3"] = "Mês de referência (1 a 12)", mes
    ws["A4"], ws["B4"] = "Ano", ano
    ws["B3"].fill = PatternFill("solid", fgColor=ENTRADA)
    kpis = [
        ("META 2 · POSITIVAÇÃO", None, None),
        (f"Pedidos {base_ano} no mês", f"=INDEX('M2 Meta x Realizado'!B4:B15,B3)", "0"),
        (f"Meta {ano} no mês", f"=INDEX('M2 Meta x Realizado'!C4:C15,B3)", "0.0"),
        (f"Realizado {ano} no mês", f"=INDEX('M2 Meta x Realizado'!D4:D15,B3)", "0"),
        ("Meta acumulada até o mês", f"=INDEX('M2 Meta x Realizado'!F4:F15,B3)", "0.0"),
        ("Realizado acumulado até o mês", f"=INDEX('M2 Meta x Realizado'!G4:G15,B3)", "0"),
        ("% da meta acumulada", f"=INDEX('M2 Meta x Realizado'!H4:H15,B3)", "0%"),
        ("Clientes previstos que não recompraram", "=COUNTIF('M2 Acompanhamento'!F:F,\"Não recomprou\")", "0"),
        ("Clientes previstos ainda sem lançamento", "=COUNTIF('M2 Acompanhamento'!F:F,\"Previsto (sem lançamento)\")", "0"),
        ("META 3a · S.A.", None, None),
        ("S.A. registradas", "=COUNTA('M3a SA'!A:A)-1", "0"),
        ("S.A. de processo de exportação", "=COUNTIF('M3a SA'!G:G,\"Processo de exportação\")", "0"),
        ("Concluídas no prazo", "=COUNTIF('M3a SA'!L:L,\"No prazo\")", "0"),
        ("Vencidas sem baixa (confirmar conclusão)", "=COUNTIF('M3a SA'!L:L,\"Vencida sem baixa\")", "0"),
        ("Cobranças registradas", "=SUM('M3a SA'!N:N)", "0"),
        ("META 3b · CUSTOS", None, None),
        ("Economia / custo evitado no ano (R$)", "=SUM('M3b Ações'!G:G)", "#,##0.00"),
        ("META 3c · EMBALAGENS", None, None),
        ("Reaproveitado no mês (R$)", "=INDEX('M3c Resumo'!B4:B15,B3)", "#,##0.00"),
        ("Reaproveitado no ano (R$)", "=INDEX('M3c Resumo'!D4:D15,B3)", "#,##0.00"),
        ("PENDÊNCIAS", None, None),
        ("Itens na aba Pendências", "=COUNTA(Pendências!B:B)-1", "0"),
    ]
    for i, (rot_, f_, fmt) in enumerate(kpis, start=6):
        c1 = ws.cell(row=i, column=1, value=rot_)
        if f_ is None:
            c1.font = Font(bold=True, color="FFFFFF")
            c1.fill = PatternFill("solid", fgColor=AZUL)
            ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=AZUL)
            continue
        c2 = ws.cell(row=i, column=2, value=f_)
        c2.number_format = fmt
        c2.fill = PatternFill("solid", fgColor=CALC)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    mestre_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(mestre_path)

    # ================================================================ relatório
    doc = Doc()
    doc.h(f"PPR Comércio Exterior · Relatório de {nome_mes}")
    doc.p(f"Fontes: {', '.join(sorted({f.name.split('-', 1)[-1] for v in arq.values() for f in v}))}. "
          f"Base da positivação: {base_ano}.")

    doc.h("1. Resumo", 2)
    acum_meta = meta_mes[meta_mes.index <= mes].sum()
    acum_real = real_mes[real_mes.index <= mes].sum()
    resumo = [
        ["1 · Volume (75 t/mês)", "Apurado pela empresa nos pedidos lançados.", "Sem entrega"],
        ["2 · Positivação +20%",
         f"Meta {MESES_TXT[mes - 1]}: {n(meta_mes[mes])} pedidos ({n(base.sum()[mes], 0)} em {base_ano}). Meta acumulada "
         f"jan-{MESES_TXT[mes - 1].lower()}: {n(acum_meta)}. " +
         (f"Realizado: {n(real_mes[mes], 0)} no mês, {n(acum_real, 0)} no acumulado ({pct(acum_real / acum_meta if acum_meta else None)})."
          if tem_realizado else f"Realizado {ano}: PENDENTE (falta lançar os pedidos de {ano})."),
         "Parcial" if not tem_realizado else ("Na meta" if acum_real >= acum_meta else "Abaixo da meta")],
        ["3a · Gestão das S.A.",
         f"{len(unicas)} S.A. abertas no mês ({int((unicas.natureza == 'Processo de exportação').sum()) if len(unicas) else 0} ligadas "
         "ao processo de exportação). Prazo por setor: PENDENTE (exportação sem data de conclusão)." if len(unicas) else "Sem S.A.",
         "Parcial"],
        ["3b · Redução de custo",
         f"Custo evitado com embalagem reaproveitada: {rs(uso_mes.valor.sum() if len(uso_mes) else 0)} no mês, "
         f"{rs(uso[uso.mes <= mes].valor.sum() if len(uso) else 0)} no ano. Demais ações: valores PENDENTES.", "Parcial"],
        ["3c · Giro de embalagens",
         f"{rs(uso_mes.valor.sum() if len(uso_mes) else 0)} reaproveitados em {uso_mes.invoice.nunique() if len(uso_mes) else 0} "
         f"pedido(s) no mês; {rs(uso[uso.mes <= mes].valor.sum() if len(uso) else 0)} no ano.", "Entregue"],
    ]
    doc.tabela(pd.DataFrame(resumo, columns=["Meta", "Resultado", "Status"]))

    # --- meta 2
    doc.h("2. Meta 2 · Positivação de clientes", 2)
    doc.p(f"A meta é repetir, mês a mês, o número de pedidos de {base_ano} com 20% a mais. Em {base_ano} foram "
          f"{n(base.values.sum(), 0)} pedidos de {pos_2025} clientes (de {len(clientes_base)} cadastrados); a meta de {ano} "
          f"soma {n(meta_mes.sum())} pedidos.")
    t = pd.DataFrame({"Mês": MESES_TXT, f"{base_ano}": [n(v, 0) for v in base.sum().values],
                      f"Meta {ano}": [n(v) for v in meta_mes.values],
                      f"Realizado {ano}": [(n(real_mes[m_], 0) if tem_realizado and m_ <= mes else ("PENDENTE" if m_ <= mes else ""))
                                          for m_ in range(1, 13)],
                      "Clientes que compraram em " + str(base_ano): [", ".join(sorted(c for c in clientes_base
                                                                                     if base.loc[c, m_] > 0)) for m_ in range(1, 13)]})
    doc.tabela(t)
    doc.p(f"Concentração: {top.index[0]} respondeu por {n(top.iloc[0], 0)} dos {n(top.sum(), 0)} pedidos de {base_ano} "
          f"({pct(concentr)}). Os meses com meta mais alta ({', '.join(MESES_TXT[m_ - 1] for m_ in meta_mes.nlargest(3).index)}) "
          f"dependem dele. Qualquer pedido a menos desse cliente precisa ser compensado por outros clientes.")
    if sem_compra:
        doc.p(f"Clientes cadastrados sem nenhum pedido em {base_ano} ({len(sem_compra)}): "
              + ", ".join(c for c in sem_compra) + ". Cada um que voltar a comprar é positivação acima da base.")
    if novos_conf:
        doc.p(f"Clientes novos em {ano} (confirmados pelo gerente; fora da base de {base_ano}): {', '.join(novos_conf)}. "
              "Cada pedido deles é positivação acima da base e soma para os +20%.")
    if novos_duv:
        doc.p(f"Clientes com embarque em {ano} que não estão na base de {base_ano}: {', '.join(novos_duv)}. "
              "Se forem novos, somam para os +20%.")
    if evid:
        doc.h(f"2.1 Embarques de {ano} já identificados nos arquivos (para conferir e lançar)", 3)
        doc.tabela(pd.DataFrame([{"Cliente": c, "Meses e invoices": "; ".join(
            f"{MESES_TXT[k - 1]}: {', '.join(sorted(v))}" for k, v in sorted(ev.items()))} for c, ev in sorted(evid.items())]))
        doc.p("Esta lista vem dos arquivos de embalagem e não é a contagem oficial de pedidos. Serve para começar o "
              f"lançamento na aba 'M2 Realizado {ano}'.")
    doc.h(f"2.2 Clientes previstos para {MESES_EXT[mes - 1]} (compraram no mesmo mês de {base_ano})", 3)
    if esperados:
        doc.tabela(pd.DataFrame([{"Cliente": c, f"Pedidos em {MESES_TXT[mes - 1]}/{base_ano}": n(base.loc[c, mes], 0),
                                  "Frequência": (cad_b.loc[c, "frequencia_txt"] if c in cad_b.index else ""),
                                  f"Pedidos em {MESES_TXT[mes - 1]}/{ano}": (n(realizado.get((c, mes), 0), 0) if tem_realizado else "PENDENTE"),
                                  "Motivo se não comprou": next((str(r[3]) for r in linhas_mot if str(r[0]) == rot and r[1] == c and pd.notna(r[3]) and r[3]), "PENDENTE")}
                                 for c in esperados]))
    doc.h("2.3 Agenda dos próximos meses (contatar antes do pedido cair)", 3)
    doc.tabela(pd.DataFrame([{"Mês": MESES_TXT[m_ - 1], f"Meta {ano}": n(meta_mes[m_]),
                              "Clientes que compraram nesse mês em " + str(base_ano): ", ".join(
                                  f"{c} ({n(base.loc[c, m_], 0)})" for c in cs)} for m_, cs in prox.items()]))
    nao_cump = list(cad_b[cad_b.cumprido.astype(str).str.upper().str.startswith("N")].index) if "cumprido" in cad_b else []
    mot_base = {r[1]: r for r in linhas_mot if str(r[0]) == str(base_ano)}
    if nao_cump:
        doc.h(f"2.4 Clientes que não cumpriram a periodicidade em {base_ano}", 3)
        doc.tabela(pd.DataFrame([{"Cliente": c, "Frequência": cad_b.loc[c, "frequencia_txt"],
                                  f"Pedidos {base_ano}": n(base.loc[c].sum(), 0) if c in base.index else "",
                                  "Motivo": (mot_base[c][3] if c in mot_base and pd.notna(mot_base[c][3]) else "PENDENTE"),
                                  "Detalhe": (mot_base[c][4] if c in mot_base and pd.notna(mot_base[c][4]) else ""),
                                  "Ação": (mot_base[c][5] if c in mot_base and pd.notna(mot_base[c][5]) else "")}
                                 for c in nao_cump]))
        faltam = [c for c in nao_cump if c not in mot_base]
        if faltam:
            pend.append(f"Meta 2: motivo de não cumprimento da periodicidade em {base_ano} ({', '.join(faltam)}).")

    # --- meta 3a
    doc.h("3. Meta 3a · Gestão das solicitações de atividade (S.A.)", 2)
    if len(unicas):
        doc.p(f"{len(unicas)} S.A. abertas em {nome_mes} (a exportação lista {len(sa_mes)} linhas; as repetidas e o "
              "mesmo código em duas abas foram contados uma vez). Separei as que mexem no prazo do cliente das que são apoio interno, porque só as primeiras "
              "entram na conta do prazo de embarque.")
        g = unicas.groupby(["setor", "natureza"]).size().unstack(fill_value=0).reset_index()
        g.columns.name = None
        doc.tabela(g)
        proc = unicas[unicas.natureza == "Processo de exportação"]
        if len(proc):
            doc.h("3.1 S.A. ligadas ao processo de exportação", 3)
            doc.tabela(pd.DataFrame({"Código": proc.codigo, "Data": proc.data_criacao.map(dt), "Setor": proc.setor,
                                     "Tipo": proc.tipo, "Descrição": proc.descricao.str.slice(0, 90),
                                     "Pedidos citados": proc.pedidos_citados, "Prazo sugerido (dias)": proc.prazo,
                                     "Conclusão": proc.data_conclusao.map(dt).replace("", "PENDENTE")}))
        canc = unicas[unicas.tipo.str.contains("CANCELAMENTO DE PEDIDO|ALTERAÇÃO DE PEDIDO", regex=True)]
        if len(canc):
            peds_c = sorted({p for s_ in canc.pedidos_citados for p in s_.split(", ") if p})
            doc.p(f"Pedidos cancelados ou com data de entrega alterada no mês: {', '.join(peds_c)}.")
            doc.tabela(pd.DataFrame({"S.A.": canc.codigo, "Tipo": canc.tipo, "Pedidos": canc.pedidos_citados,
                                     "Motivo": canc.obs.fillna("PENDENTE")}))
            if canc.obs.isna().any():
                pend.append(f"Meta 3a/2: motivo dos pedidos cancelados ou com entrega alterada ({', '.join(peds_c)}).")
            elif canc.obs.astype(str).str.contains("ajuste|corre", case=False).any():
                doc.p("Os cancelamentos foram feitos para ajuste ou correção do pedido, não por perda de venda. Se o pedido "
                      "foi reemitido, ele continua contando na positivação; o cancelado não deve ser contado em dobro.")
        if len(usuarios):
            doc.p("Solicitações por usuário (aba de percentuais da exportação): " + "; ".join(
                f"{r.usuario.title()} {n(r.solicitacoes, 0)}" for r in usuarios.itertuples()) + ".")
    if len(fluxo):
        doc.h(f"3.2 Referência de prazo: processos de {base_ano} no Controle de Fluxo", 3)
        ok = fluxo[fluxo.dias_total.notna()]
        etapas = [c for c in fluxo.columns if c not in ("invoice", "pedido", "cliente", "pais", "data_pedido", "previsao_saida",
                                                         "saida_navio", "tinha_estoque", "status", "motivo", "dias_total",
                                                         "dias_vs_previsao", "data_pedido_suspeita")]
        med = pd.DataFrame([{"Etapa": e, "Processos medidos": int(fluxo[e].notna().sum()), "Média (dias)": n(fluxo[e].mean()),
                             "Máximo (dias)": n(fluxo[e].max(), 0)} for e in etapas if fluxo[e].notna().any()])
        doc.p(f"Só {len(fluxo)} processos foram registrados (até {dt(fluxo.data_pedido[~fluxo.data_pedido_suspeita].max())[3:]}), "
              f"todos com produto em estoque. Nos {len(ok)} com data de saída do navio, o pedido levou em média "
              f"{n(ok.dias_total.mean(), 0)} dias até embarcar (prazo combinado: 30). O motivo mais anotado foi demora na produção.")
        doc.tabela(med)
        doc.p("As etapas de compra de embalagem, arte, laboratório e PCP estão vazias. São justamente as que dependem de "
              "S.A. a outros setores, e é onde a cobrança faz diferença.")
    doc.h("3.3 Rotina de cobrança proposta (passa a gerar a evidência do relatório)", 3)
    doc.lista([
        "Dia 0: ao lançar o pedido, abrir as S.A. já com a data limite de cada setor (aba 'M3a Prazos'), contada de trás para "
        "frente a partir do embarque em 30 dias.",
        "Data limite menos 2 dias: lembrete ao setor (e-mail curto com pedido, atividade e data de embarque do cliente).",
        "Data limite: cobrança formal; anotar na aba 'M3a SA' (Nº cobranças e Última cobrança).",
        "Data limite mais 3 dias: escalar ao gestor do setor com cópia ao seu gestor.",
        "Fim do mês: o relatório sai sozinho da aba 'M3a SA' (% no prazo por setor, atraso médio e cobranças feitas).",
    ])

    # --- meta 3b
    doc.h("4. Meta 3b · Redução de custos", 2)
    acoes_mes = [r for r in acoes if str(r[0]) == rot]
    if acoes_mes:
        doc.tabela(pd.DataFrame([{"Ação": r[2][:110], "Tipo de ganho": r[3], "Valor (R$)": rs((r[4] or 0) - (r[5] or 0)),
                                  "Evidência": r[7]} for r in acoes_mes]))
    doc.p("Custo evitado = valor de embalagem que estava parada e foi usada no pedido, em vez de comprar embalagem nova "
          "ou dar baixa no estoque. O mesmo valor aparece na meta 3c; no relatório consolidado, contar uma vez só.")
    if len(gastos):
        cats = ["Taxas locais", "Desembaraço aduaneiro", "Estufagem do container", "Documentos", "Custos no porto",
                "Etiquetas", "Courier", "Comissão de agente", "Outros", "Diferença de frete"]
        cats = [c for c in cats if c in gastos.columns]
        tot_c = gastos[cats].sum().sort_values(ascending=False)
        tot_c = tot_c[tot_c > 0]
        doc.h(f"4.1 Onde está o custo por processo (Controle de Gastos, {len(gastos)} processos de {base_ano})", 3)
        doc.tabela(pd.DataFrame({"Item": tot_c.index, "Total (R$)": [rs(v) for v in tot_c.values],
                                 "% do total": [pct(v / tot_c.sum()) for v in tot_c.values]}))
        doc.p(f"Total de {rs(tot_c.sum())} para {n(gastos['Unidades de produto vendidas'].sum(), 0)} unidades vendidas. "
              f"{tot_c.index[0]} e {tot_c.index[1].lower()} somam {pct(tot_c.iloc[:2].sum() / tot_c.sum())}. São os primeiros "
              "pontos a negociar.")
    if len(cambio):
        c_ok = cambio[~cambio.duplicada]
        doc.p(f"Faturado x câmbio {base_ano}: {rs(c_ok.valor_nf.sum())} faturados e {rs(c_ok.recebido_cambio.sum())} recebidos "
              f"({rs(c_ok.variacao.sum())}, {pct(c_ok.variacao.sum() / c_ok.valor_nf.sum())}), sem as linhas repetidas. "
              "Não é redução de custo, mas mostra quanto a variação cambial pesa no resultado.")
    doc.h("4.2 Ações para medir a partir de agora (falta só o valor)", 3)
    candidatas = ["Amostras: separar as enviadas por conta da empresa das pagas pelo cliente (courier por conta do cliente = "
                  "custo transferido)."]
    if len(unicas) and unicas.tipo.str.contains("AMOSTRA").any():
        candidatas[0] += f" No mês houve {int(unicas.tipo.str.contains('AMOSTRA').sum())} S.A. de amostras."
    if len(unicas) and unicas.tipo.str.contains("MOTORISTA").any():
        candidatas.append(f"Deslocamentos: {int(unicas.tipo.str.contains('MOTORISTA').sum())} S.A. de motorista no mês "
                          "(levar e buscar documento no cartório). Juntar numa ida só e registrar o custo de cada ida.")
    candidatas += ["Comissão de agente: maior item de custo por processo; revisar percentual nos próximos pedidos.",
                   "Estufagem e taxas locais: cotar com mais de um operador e registrar a cotação perdedora como referência.",
                   "Documentos e courier: enviar documentos digitais quando o país aceitar; cada envio de courier evitado é economia.",
                   "Inspeções: agendar no mesmo dia para pedidos que embarcam juntos."]
    doc.lista(candidatas)
    if not [r for r in acoes_mes if r[8] != "Automático"]:
        doc.p("Sem outras ações de custo com valor registradas no mês. Quando houver, basta informar a ação, quanto "
              "custaria e quanto custou; o agente lança na aba 'M3b Ações'.")

    # --- meta 3c
    doc.h("5. Meta 3c · Giro de embalagens", 2)
    if len(uso):
        rm = uso[uso.mes <= mes].groupby("mes").agg(valor=("valor", "sum"), pedidos=("invoice", "nunique"),
                                                     clientes=("cliente", lambda s: ", ".join(sorted(set(s)))))
        doc.tabela(pd.DataFrame({"Mês": [MESES_TXT[m_ - 1] for m_ in rm.index], "Valor reaproveitado": [rs(v) for v in rm.valor],
                                 "Pedidos": rm.pedidos.values, "Clientes": rm.clientes.values}))
        doc.p(f"Total de {ano} até {MESES_TXT[mes - 1].lower()}: {rs(rm.valor.sum())}.")
    if len(uso_mes):
        doc.h(f"5.1 Detalhe de {MESES_EXT[mes - 1]}", 3)
        doc.tabela(pd.DataFrame({"Cliente": uso_mes.cliente, "Invoice": uso_mes.invoice, "Produto": uso_mes.descricao,
                                 "Código": uso_mes.codigo_embalagem, "Tipo": uso_mes.tipo.fillna("").str.strip(),
                                 "Qtd": uso_mes.qtd_embalagem.map(lambda v: n(v, 0) if pd.notna(v) else ""),
                                 "Valor": uso_mes.valor.map(rs),
                                 "Dias parada": uso_mes.dias_parada.map(lambda v: n(v, 0) if pd.notna(v) else "")}))
        velho = uso_mes.sort_values("dias_parada", ascending=False).iloc[0]
        doc.p(f"A embalagem mais antiga usada no mês estava parada há {n(velho.dias_parada, 0)} dias "
              f"({velho.descricao}, cód. {velho.codigo_embalagem}).")
    if len(vp):
        doc.h("5.2 Estoque de embalagens sem movimentação", 3)
        tot_ = vp[vp.bloco.str.startswith("Descrição Estoque Total")]
        pv = tot_.pivot_table(index="faixa", columns="data", values="valor", aggfunc="sum", sort=False)
        doc.tabela(pd.DataFrame({"Faixa": pv.index, **{d.strftime("%m/%Y"): [rs(v) for v in pv[d]] for d in pv.columns}}))
        mais90 = vp[vp.bloco.str.contains("90 dias")]
        if len(mais90):
            s_ = mais90.sort_values("data")
            doc.p(f"Estoque parado há mais de 90 dias: {rs(s_.valor.iloc[0])} em {s_.data.iloc[0].strftime('%m/%Y')} e "
                  f"{rs(s_.valor.iloc[-1])} em {s_.data.iloc[-1].strftime('%m/%Y')} "
                  f"({pct(s_.valor.iloc[-1] / s_.valor.iloc[0] - 1)}).")
    if len(est_cli):
        doc.h("5.3 Estoque de embalagem dedicada por cliente", 3)
        for cli, g in est_cli.groupby("cliente"):
            ms = sorted(g.mes.unique())
            ini_, fim_ = ms[0], ms[-1]
            pv = g.groupby(["codigo", "descricao", "mes"]).estoque.sum().unstack()
            linhas = []
            for (cod, desc), row in pv.iterrows():
                a0, a1 = row.get(ini_), row.get(fim_)
                linhas.append({"Código": cod, "Descrição": desc, MESES_TXT[ini_ - 1]: n(a0, 0) if pd.notna(a0) else "",
                               MESES_TXT[fim_ - 1]: n(a1, 0) if pd.notna(a1) else "",
                               "Variação": n(a1 - a0, 0) if pd.notna(a0) and pd.notna(a1) else ""})
            doc.p(f"{cli}: estoque por código de {MESES_TXT[ini_ - 1]} a {MESES_TXT[fim_ - 1]}.")
            doc.tabela(pd.DataFrame(linhas))
    if pref["descritivo_embalagens"]:
        doc.p("Descritivo para a produção: com a lista de pedidos em carteira (cliente, produto, idioma, quantidade) e o "
              "estoque parado por código, o agente monta a sugestão pedido a pedido (mesmo idioma, bonificação, caixa de reforço).")
        pend.append("Meta 3c: lista de pedidos em carteira para os próximos meses (cliente, produto, idioma, quantidade) "
                    "para o descritivo de embalagens da produção.")

    # --- dados com problema e pendências
    doc.h("6. Problemas encontrados nas planilhas de origem", 2)
    doc.lista(alertas or ["Nenhum."])
    if notas:
        doc.h("7. Observações", 2)
        doc.lista(notas)
    doc.h(f"{8 if notas else 7}. O que falta para fechar o mês", 2)
    doc.lista(pend or ["Nada pendente."])

    nome = f"PPR_Comex_{a.mes}"
    doc.docx(saida / f"{nome}.docx")
    (saida / f"{nome}.md").write_text(doc.md(), encoding="utf-8")
    (saida / "Pendencias.md").write_text(
        f"# Pendências · {nome_mes}\n\n" + "\n".join(f"{i}. {t}" for i, t in enumerate(pend, 1)) +
        "\n\n## Corrigir nas planilhas de origem\n\n" + "\n".join(f"- {t}" for t in alertas) + "\n", encoding="utf-8")
    resumo_txt = [
        f"PPR Comex · {nome_mes}",
        f"Positivação: meta de {n(meta_mes[mes])} pedidos no mês e {n(acum_meta)} no acumulado; " +
        (f"realizado {n(real_mes[mes], 0)} / {n(acum_real, 0)}." if tem_realizado else "realizado em lançamento."),
        f"S.A.: {len(unicas)} abertas no mês, {int((unicas.natureza == 'Processo de exportação').sum()) if len(unicas) else 0} "
        "ligadas a pedidos; controle de prazo e cobrança por setor implantado a partir deste mês.",
        f"Custos: {rs(uso_mes.valor.sum() if len(uso_mes) else 0)} de custo evitado com embalagem reaproveitada; demais ações em levantamento.",
        f"Embalagens: {rs(uso_mes.valor.sum() if len(uso_mes) else 0)} reaproveitados no mês, "
        f"{rs(uso[uso.mes <= mes].valor.sum() if len(uso) else 0)} no ano.",
    ]
    (saida / "Resumo_gestor.txt").write_text("\n".join(resumo_txt) + "\n", encoding="utf-8")

    print(mestre_path)
    print(saida)
    for t in pend:
        print("PENDENTE:", t)
    for t in alertas:
        print("ALERTA:", t)


if __name__ == "__main__":
    main()
